"""KUK 3.3 — Laporan hasil pengujian didokumentasikan berdasarkan data pengujian.

Tujuh bagian (materi slide 43): (1) identitas produk, versi, ruang lingkup; (2) lingkungan & perangkat uji;
(3) metode & parameter (UK 1); (4) rekap TC per status & tingkat lulus; (5) temuan & ketidaksesuaian terhadap standar;
(6) bukti: log, tangkapan layar, hash; (7) kesimpulan & tanda tangan penguji.
Ekspor PDF (ReportLab) dan DOCX (python-docx), header tabel #3D7A74, "Halaman X dari Y", plus uk3_hasil_uji.json untuk UK-4.
Semua isi dibaca dari berkas run (run.json, analisis.json, kesimpulan.json, prep.json, manifest)."""
import datetime as dt
import json
import shutil
from pathlib import Path

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.platypus import CondPageBreak, NextPageTemplate, PageBreak, Spacer

from std_report.report.template import (Caption, Heading, ListOf, NumberedCanvas, P, ReportDoc, cm, data_table, kv_table,
                                        make_toc, mono)

from .. import CATEGORIES, KUK, UNIT, UNIT_TITLE, __version__, rel
from .. import products as Pr
from ..evidence import append, verify
from ..prep.documents import latest as latest_prep
from ..run.planner import STAGES
from ..run.runner import load_run, run_dir, runs_index, uk3_dir

ACCENT = "#3D7A74"
ST_COLOR = {"LULUS": "#2E7D32", "GAGAL": "#C62828", "TERBLOKIR": "#B26A00", "TIDAK DAPAT DIUJI": "#6B7280"}
CAT_COLOR = {"Memenuhi": "#2E7D32", "Memenuhi dengan Catatan": "#3D7A74", "Inkonklusif": "#B26A00", "Tidak Memenuhi": "#C62828"}
PW = A4[0] - 4 * cm
SECTIONS = ["Identitas produk, versi, dan ruang lingkup", "Lingkungan & perangkat uji", "Metode & parameter uji (UK 1)",
            "Rekap test case", "Temuan & ketidaksesuaian terhadap standar", "Bukti: log, tangkapan layar, hash", "Kesimpulan & tanda tangan penguji"]
BULAN = ["Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli", "Agustus", "September", "Oktober", "November", "Desember"]


def _tgl(iso):
    d = dt.datetime.fromisoformat(iso)
    return f"{d.day} {BULAN[d.month - 1]} {d.year}"


def gather(product_id, run_id, out_root=None) -> dict:
    """Kumpulkan semua data laporan; analisis & kesimpulan dibuat bila belum ada."""
    d = run_dir(product_id, run_id, out_root)
    if not (d / "analisis.json").exists():
        from ..analysis.analyze import analyze
        analyze(product_id, run_id, out_root)
    if not (d / "kesimpulan.json").exists():
        from ..conclude.decide import conclude
        conclude(product_id, run_id, out_root)
    rd = lambda n: json.loads((d / n).read_text(encoding="utf-8"))  # noqa: E731
    return {"run": load_run(product_id, run_id, out_root), "analisis": rd("analisis.json"), "kesimpulan": rd("kesimpulan.json"),
            "prep": latest_prep(product_id, out_root), "manifest": rd("evidence_manifest.json"), "verifikasi": verify(d),
            "riwayat": runs_index(product_id, out_root), "produk": Pr.REGISTRY[product_id], "dir": d,
            "screenshots": sorted((d / "evidence" / "screenshots").glob("*.png")) if (d / "evidence" / "screenshots").exists() else []}


def hasil_uji_json(g) -> dict:
    """Keluaran untuk UK-4 (rekomendasi perbaikan)."""
    run, k = g["run"], g["kesimpulan"]
    return {"schema": "cryptan.uk3.hasil_uji.v1", "app": f"Cryptan.ID Test Suite · uk3_pengujian {__version__}", "unit": UNIT,
            "produk": run["produk"], "nama_produk": run["nama_produk"], "run_id": run["run_id"], "waktu": run["selesai"],
            "versi": {v: {"sha256": x["sha256"], "label": x["label"], "rekap": k["per_versi"][v]["rekap"],
                          "tingkat_lulus": k["per_versi"][v]["tingkat_lulus"], "keseluruhan": k["per_versi"][v]["keseluruhan"],
                          "sasaran": {s: {"nama": o["nama"], "kategori": o["kategori"], "alasan": o["alasan"]}
                                      for s, o in k["per_versi"][v]["sasaran"].items()},
                          "tc": [{f: t.get(f) for f in ("id", "judul", "jenis", "tahap", "status", "actual", "bukti")} for t in x["tc"]]}
                      for v, x in run["versi"].items()},
            "temuan": [{f: t.get(f) for f in ("id", "judul", "versi", "test_case", "dampak", "acuan", "keparahan", "rekomendasi")} |
                       {"cvss_vector": (t.get("cvss") or {}).get("vector"), "cvss_score": (t.get("cvss") or {}).get("base_score")}
                       for t in k["temuan"]],
            "pernyataan": k["pernyataan"], "deviasi": run["deviasi"], "evidence_manifest_ok": g["verifikasi"]["ok"],
            "handoff_uk4": {"note": "Masukan UK-4: setiap temuan → rekomendasi perbaikan & prioritas berdasarkan CVSS",
                            "temuan": len(k["temuan"])}}


# ----------------------------------------------------------------------------------------------- PDF
def _badge_rows(rows, col, palette):
    """Latar warna status + teks status putih tebal (status selalu tertulis, bukan hanya warna)."""
    cmds = []
    for i, r in enumerate(rows, 1):
        st = r[col] if isinstance(r[col], str) else ""
        c = palette.get(st)
        if c:
            cmds += [("BACKGROUND", (col, i), (col, i), colors.HexColor(c)), ("VALIGN", (col, i), (col, i), "MIDDLE")]
            r[col] = P(f"<b>{st}</b>", "cell_badge", raw=True)
    return cmds


def _vt(inv):
    rows = [[c["alat"], c["uji"], mono(c["expected"], "cell_mono_small"), mono(c["actual"], "cell_mono_small"),
             "LULUS" if c["lulus"] else "GAGAL"] for c in inv.get("verifikasi", [])]
    ex = _badge_rows(rows, 4, ST_COLOR)
    return data_table(["Alat", "Uji", "Expected", "Aktual", "Hasil"], rows, [2.8 * cm, 3.6 * cm, 4.2 * cm, 4.2 * cm, PW - 14.8 * cm], extra=ex)


def _white(text):
    return P(f"<b>{text}</b>", "cell", raw=True)


def build_pdf(g, path, penguji=None) -> Path:
    run, k, prep, an = g["run"], g["kesimpulan"], g["prep"], g["analisis"]
    penguji = penguji or (prep or {}).get("ruang_lingkup", {}).get("penguji") or "Penguji"
    story, nt = [], [0]

    def cap(text):
        nt[0] += 1
        return Caption("table", nt[0], text)

    def h(i, key):
        return Heading(f"{i}  {SECTIONS[i - 1]}", 0, key)

    vers = list(run["versi"])
    # sampul
    story += [Spacer(1, 2 * cm), P("LAPORAN HASIL PENGUJIAN PRODUK KRIPTOGRAFI", "cover_sub"), Spacer(1, 0.3 * cm),
              P(f"{run['nama_produk']} — versi {', '.join(vers)}", "cover_title"), Spacer(1, 0.4 * cm),
              P(f"Unit kompetensi {UNIT} · {UNIT_TITLE}", "cover_sub"), Spacer(1, 1 * cm),
              kv_table([("Penguji", penguji), ("Produk", f"{run['nama_produk']} ({run['produk']})"),
                        ("Versi & sidik jari", "; ".join(f"v{v} · {x['sidik_jari']} · {x['label']}" for v, x in run["versi"].items())),
                        ("Run ID", run["run_id"]), ("Waktu uji", f"{run['mulai']} – {run['selesai']}"),
                        ("Alasan run", run["alasan_run"]), ("Tanggal laporan", _tgl(dt.datetime.now().astimezone().isoformat())),
                        ("Aplikasi", f"Cryptan.ID Test Suite · uk3_pengujian {__version__}")], [4 * cm, PW - 4 * cm]),
              Spacer(1, 1.5 * cm), P("© Cryptan.ID — made by Dea Saka Kurnia Putra", "small"),
              NextPageTemplate("portrait"), PageBreak(), P("Daftar Isi", "h1"), make_toc(), Spacer(1, 0.5 * cm),
              P("Daftar Tabel", "h1"), ListOf("LOT"), PageBreak()]
    # 1 identitas
    spec = g["produk"].get("spesifikasi", {})
    story += [h(1, "s1"), P(f"{g['produk']['nama']} — {g['produk']['jenis']}. Pemilik: {g['produk']['pemilik']}."),
              cap("Spesifikasi desain produk"), kv_table([(a, b) for a, b in spec.items()], [3.5 * cm, PW - 3.5 * cm]), Spacer(1, 0.3 * cm),
              cap("Versi produk yang diuji"),
              data_table(["Versi", "Berkas", "SHA-256", "Label"], [[f"v{v}", x["berkas"] or "—", mono(x["sha256"] or "—", "cell_mono_small"),
                                                                   x["label"]] for v, x in run["versi"].items()],
                         [1.4 * cm, 4.6 * cm, 7.6 * cm, PW - 13.6 * cm])]
    sc = (prep or {}).get("ruang_lingkup", {})
    story += [Spacer(1, 0.3 * cm), cap("Ruang lingkup, batasan, dan izin"),
              kv_table([("Izin tertulis", f"{sc.get('izin', {}).get('nomor', '—')} oleh {sc.get('izin', {}).get('oleh', '—')}, "
                                          f"{sc.get('izin', {}).get('tanggal', '—')}"),
                        ("Jadwal", f"{sc.get('jadwal', {}).get('mulai', '—')} s.d. {sc.get('jadwal', {}).get('selesai', '—')}"),
                        ("Batasan", "; ".join(sc.get("batasan", []))), ("Kerahasiaan", sc.get("kerahasiaan", "—"))], [3.5 * cm, PW - 3.5 * cm]),
              PageBreak()]
    # 2 lingkungan
    inv = (prep or {}).get("inventaris", {})
    story += [h(2, "s2"), P(run["catatan_lingkungan"], "cell_mono"), Spacer(1, 0.2 * cm), cap("Perangkat pendukung (KUK 1.2)"),
              data_table(["Perangkat", "Versi", "Wajib", "Dipakai untuk"],
                         [[r["perangkat"], r["versi"] or "—", "ya" if r["wajib"] else "opsional", r["dipakai_untuk"]] for r in inv.get("perangkat", [])],
                         [4 * cm, 5 * cm, 1.6 * cm, PW - 10.6 * cm]),
              Spacer(1, 0.3 * cm), cap("Verifikasi alat dengan nilai yang diketahui"),
              _vt(inv),
              Spacer(1, 0.3 * cm), cap("Checklist kesiapan (entry criteria)"),
              data_table(["Butir", "Status", "Keterangan"], [[i["uraian"], i["status"], i["kekurangan"] or "—"]
                                                            for i in (prep or {}).get("checklist", {}).get("items", [])],
                         [7 * cm, 1.6 * cm, PW - 8.6 * cm]), PageBreak()]
    # 3 metode
    story += [h(3, "s3"), P("Metode analisis dan aturan keputusan per metode (materi T6–T7). Expected value berasal dari skenario "
                            "(gaya UK-2) atau vektor resmi, tidak dari produk yang diuji."), cap("Metode dan aturan keputusan"),
              data_table(["Metode", "Aturan keputusan", "TC"], [[m["metode"], m["aturan"], ", ".join(sorted({b["tc"] for b in m["baris"]}))]
                                                                for m in an["metode"]], [3.6 * cm, 7 * cm, PW - 10.6 * cm]),
              Spacer(1, 0.3 * cm), cap("Parameter statistik dikunci sebelum uji"),
              data_table(["TC", "Parameter terkunci"], [[t, json.dumps(v, ensure_ascii=False)] for t, v in run["parameter_terkunci"].items()]
                         or [["—", "—"]], [3 * cm, PW - 3 * cm])]
    if run["deviasi"]:
        story += [Spacer(1, 0.2 * cm), cap("Deviasi parameter setelah dikunci"),
                  data_table(["TC", "Parameter", "Terkunci", "Dipakai", "Alasan"],
                             [[x["tc"], x["parameter"], str(x["terkunci"]), str(x["dipakai"]), x["alasan"]] for x in run["deviasi"]],
                             [2.2 * cm, 2.4 * cm, 2.4 * cm, 2.4 * cm, PW - 9.4 * cm])]
    uk1 = [f for f in g["prep"]["dokumen"] if f["jenis"].startswith("Metode")] if prep else []
    if uk1:
        story += [P(f"Metode & parameter dari UK-1: {uk1[0]['berkas']} (SHA-256 {str(uk1[0]['sha256'])[:16]}…).", "small")]
    story += [Spacer(1, 0.3 * cm), cap("Tahapan pengujian (KUK 2.1)"),
              data_table(["Tahap", "Nama", "Alasan urutan", "TC"], [[str(s["tahap"]), s["nama"], s["alasan"], ", ".join(s["tc"])]
                                                                   for s in run["rencana"]["tahapan"]], [1.2 * cm, 3.6 * cm, 4.6 * cm, PW - 9.4 * cm])]
    for w in run["rencana"]["peringatan"]:
        story.append(P(f"Peringatan: {w}", "small"))
    story.append(PageBreak())
    # 4 rekap
    rk = [[f"v{v}", *[str(x["rekap"][s]) for s in ("LULUS", "GAGAL", "TERBLOKIR", "TIDAK DAPAT DIUJI")], str(x["total"]),
           f"{x['rekap']['LULUS']}/{x['total']} = {x['tingkat_lulus'] * 100:.1f}%".replace(".", ",")] for v, x in k["per_versi"].items()]
    story += [h(4, "s4"), cap("Rekap test case per status dan tingkat lulus"),
              data_table(["Versi", "LULUS", "GAGAL", "TERBLOKIR", "TDD", "Total", "Tingkat lulus"], rk,
                         [1.6 * cm, 1.7 * cm, 1.7 * cm, 2.1 * cm, 1.5 * cm, 1.5 * cm, PW - 10.1 * cm]),
              Spacer(1, 0.3 * cm), cap("Lembar hasil berdampingan")]
    rows, ex = [], []
    for i, tc in enumerate(run["versi"][vers[0]]["tc"]):
        sts = [run["versi"][v]["tc"][i]["status"] for v in vers]
        rows.append([mono(tc["id"]), P(tc["judul"], "cell_small"), str(tc["tahap"]), tc["jenis"]] + [P(f"<b>{s}</b>", "cell_badge", raw=True) for s in sts])
        for j, s in enumerate(sts):
            ex += [("BACKGROUND", (4 + j, i + 1), (4 + j, i + 1), colors.HexColor(ST_COLOR[s]))]
    wv = 2.2 * cm
    story += [data_table(["TC", "Judul", "Tahap", "Jenis"] + [f"v{v}" for v in vers], rows,
                         [1.9 * cm, PW - 5.8 * cm - wv * len(vers), 1.3 * cm, 2.6 * cm] + [wv] * len(vers), extra=ex),
              Spacer(1, 0.3 * cm), cap("Hasil aktual vs expected per test case")]
    det = []
    for v in vers:
        for t in run["versi"][v]["tc"]:
            e = t.get("expected") or {}
            det.append([f"v{v}", mono(t["id"]), P(str(e.get("nilai") or e.get("kriteria") or "—"), "cell_small"),
                        P(str(t["actual"]), "cell_small"), t["status"]])
    ex = _badge_rows(det, 4, ST_COLOR)
    story += [data_table(["Versi", "TC", "Expected", "Aktual", "Status"], det, [1.2 * cm, 1.9 * cm, 5.4 * cm, 6.2 * cm, PW - 14.7 * cm],
                         extra=ex), PageBreak()]
    # 5 temuan
    story += [h(5, "s5")]
    if not k["temuan"]:
        story.append(P("Tidak ada TC yang GAGAL; tidak ada temuan ketidaksesuaian pada TC yang dijalankan."))
    for f in k["temuan"]:
        story += [CondPageBreak(6 * cm), cap(f"Temuan {f['id']} — {f['judul']} (v{f['versi']})"),
                  kv_table([("ID & judul", f"{f['id']} · {f['judul']}"), ("Test case", f["test_case"]), ("Bukti", f["bukti"]),
                            ("Dampak", f["dampak"]), ("Acuan", f["acuan"] or "—"), ("Keparahan", f["keparahan"]),
                            ("Rekomendasi", f.get("rekomendasi") or "—")], [3.2 * cm, PW - 3.2 * cm])]
        if f.get("cvss"):
            story += [Spacer(1, 0.2 * cm), cap(f"Perhitungan CVSS 3.1 {f['id']}: {f['cvss']['vector']}"),
                      data_table(["Langkah", "Perhitungan", "Hasil"], [[a, b, str(c).replace(".", ",")] for a, b, c in f["cvss"]["langkah"]],
                                 [3.4 * cm, 9 * cm, PW - 12.4 * cm])]
        demo = f.get("demonstrasi_dampak")
        if demo and "berhasil" in demo:
            story += [Spacer(1, 0.2 * cm), P(f"Bukti dampak: nonce sama = {demo['nonce_sama']}; C1 ⊕ C2 = P1 ⊕ P2 = "
                                             f"{demo['c1_xor_c2_sama_dengan_p1_xor_p2']}. P2 diketahui “{demo['p2_diketahui']}” → "
                                             f"P1 dipulihkan “{demo['p1_dipulihkan']}”.", "body")]
    story.append(PageBreak())
    # 6 bukti
    vf = g["verifikasi"]
    story += [h(6, "s6"), P(f"Verifikasi manifest: {vf['kesimpulan']} ({vf['diperiksa']} berkas)."), cap("Berkas bukti dan SHA-256"),
              data_table(["Berkas", "Byte", "SHA-256"], [[mono(b["path"], "cell_mono_small"), str(b["bytes"]), mono(b["sha256"], "cell_mono_small")]
                                                         for b in g["manifest"]["berkas"]], [6.4 * cm, 1.4 * cm, PW - 7.8 * cm]),
              Spacer(1, 0.3 * cm), cap("Riwayat semua percobaan (tidak ada pengulangan diam-diam)"),
              data_table(["Run ID", "Mulai", "Alasan", "Rekap"], [[mono(x["run_id"], "cell_mono_small"), x["mulai"], x["alasan"],
                                                                  ", ".join(f"{a} {b}" for a, b in x["rekap"].items())] for x in g["riwayat"]],
                         [3.6 * cm, 3.6 * cm, 4.4 * cm, PW - 11.6 * cm])]
    if g["screenshots"]:
        from reportlab.platypus import Image
        for s in g["screenshots"]:
            story += [CondPageBreak(9 * cm), P(f"Tangkapan layar: {s.name}", "caption"), Image(str(s), width=PW, height=PW * 0.56)]
    else:
        story.append(P("Tangkapan layar: simpan berkas PNG ke evidence/screenshots/ pada folder run untuk disertakan.", "small"))
    story.append(PageBreak())
    # 7 kesimpulan
    story += [h(7, "s7")]
    for v, x in k["per_versi"].items():
        rows = [[s, o["nama"], P(f"<b>{o['kategori']}</b>", "cell_badge", raw=True), P(o["alasan"], "cell_small")] for s, o in sorted(x["sasaran"].items())]
        ex = [("BACKGROUND", (2, i), (2, i), colors.HexColor(CAT_COLOR[o["kategori"]])) for i, (_, o) in enumerate(sorted(x["sasaran"].items()), 1)]
        story += [cap(f"Kesimpulan per sasaran uji — v{v} (keseluruhan: {x['keseluruhan']})"),
                  data_table(["Sasaran", "Uraian", "Kategori", "Alasan"], rows, [1.6 * cm, 6 * cm, 3 * cm, PW - 10.6 * cm], extra=ex),
                  Spacer(1, 0.25 * cm)]
    story += [P("Rumusan kesimpulan", "h2")]
    for s, c in zip(k["pernyataan"], k["pemeriksaan_rumusan"]):
        story.append(P(f"• {s}" + ("" if c["diterima"] else " [DITOLAK pemeriksa rumusan]"), "bullet"))
    story += [P("Kesimpulan dibatasi pada versi produk, metode, dan test case yang dijalankan. TC yang TERBLOKIR atau TIDAK DAPAT "
                "DIUJI tidak dihitung sebagai LULUS. Lolos uji keacakan adalah syarat perlu, bukan syarat cukup.", "small"),
              Spacer(1, 0.4 * cm), cap("Tanda tangan penguji"),
              data_table(["Peran", "Nama", "Tanda tangan", "Tanggal"], [["Penguji", penguji, "", ""], ["Pemeriksa", "", "", ""]],
                         [3 * cm, 6 * cm, 4.5 * cm, PW - 13.5 * cm], zebra=False,
                         extra=[("TOPPADDING", (0, 1), (-1, -1), 16), ("BOTTOMPADDING", (0, 1), (-1, -1), 16)]),
              Spacer(1, 0.4 * cm), P("© Cryptan.ID — made by Dea Saka Kurnia Putra", "small")]
    meta = {"title": f"Laporan Hasil Pengujian {run['nama_produk']}", "short_title": f"Laporan Hasil Pengujian · {run['nama_produk']}",
            "penyusun": penguji, "klasifikasi": "BIASA — produk latihan", "tanggal": _tgl(dt.datetime.now().astimezone().isoformat()),
            "subject": f"{UNIT} {UNIT_TITLE}", "creator": f"Cryptan.ID Test Suite uk3_pengujian {__version__}"}
    doc = ReportDoc(path, meta)
    doc.multiBuild(story, canvasmaker=NumberedCanvas, maxPasses=6)
    return Path(path)


# ----------------------------------------------------------------------------------------------- DOCX
def build_docx(g, path, penguji=None) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    run, k, prep = g["run"], g["kesimpulan"], g["prep"]
    penguji = penguji or (prep or {}).get("ruang_lingkup", {}).get("penguji") or "Penguji"
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "DejaVu Sans", Pt(10)

    def shade(cell, hexcolor):
        tcPr = cell._tc.get_or_add_tcPr()
        s = OxmlElement("w:shd")
        s.set(qn("w:val"), "clear")
        s.set(qn("w:color"), "auto")
        s.set(qn("w:fill"), hexcolor.lstrip("#"))
        tcPr.append(s)

    def table(header, rows, status_col=None, palette=None):
        t = doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        for i, hd in enumerate(header):
            c = t.rows[0].cells[i]
            c.text = ""
            r = c.paragraphs[0].add_run(hd)
            r.bold, r.font.color.rgb, r.font.size = True, RGBColor(255, 255, 255), Pt(9)
            shade(c, ACCENT)
        hdr = t.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader")
        th.set(qn("w:val"), "true")
        hdr.append(th)
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                rr = cells[i].paragraphs[0].add_run(str(v))
                rr.font.size = Pt(8.5)
                if status_col is not None and i == status_col and str(v) in (palette or {}):
                    shade(cells[i], palette[str(v)])
                    rr.font.color.rgb, rr.bold = RGBColor(255, 255, 255), True
        doc.add_paragraph()

    def field(par, code):
        for kind, text in (("begin", None), (None, code), ("end", None)):
            r = par.add_run()
            if kind:
                e = OxmlElement("w:fldChar")
                e.set(qn("w:fldCharType"), kind)
            else:
                e = OxmlElement("w:instrText")
                e.set(qn("xml:space"), "preserve")
                e.text = text
            r._r.append(e)

    sec = doc.sections[0]
    hp = sec.header.paragraphs[0]
    hp.text = f"Laporan Hasil Pengujian · {run['nama_produk']}\tBIASA — produk latihan"
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.add_run(f"{penguji} · ")
    fp.add_run("Halaman ")
    field(fp, "PAGE")
    fp.add_run(" dari ")
    field(fp, "NUMPAGES")

    doc.add_heading("LAPORAN HASIL PENGUJIAN PRODUK KRIPTOGRAFI", 0)
    doc.add_paragraph(f"{run['nama_produk']} — versi {', '.join(run['versi'])} · {UNIT} {UNIT_TITLE}")
    table(["Butir", "Nilai"], [["Penguji", penguji], ["Run ID", run["run_id"]], ["Waktu uji", f"{run['mulai']} – {run['selesai']}"],
                               ["Alasan run", run["alasan_run"]]])
    doc.add_heading(f"1  {SECTIONS[0]}", 1)
    table(["Spesifikasi", "Nilai"], [[a, b] for a, b in g["produk"].get("spesifikasi", {}).items()])
    table(["Versi", "Berkas", "SHA-256", "Label"], [[f"v{v}", x["berkas"], x["sha256"], x["label"]] for v, x in run["versi"].items()])
    sc = (prep or {}).get("ruang_lingkup", {})
    doc.add_paragraph("Batasan: " + "; ".join(sc.get("batasan", [])))
    doc.add_heading(f"2  {SECTIONS[1]}", 1)
    doc.add_paragraph(run["catatan_lingkungan"])
    inv = (prep or {}).get("inventaris", {})
    table(["Perangkat", "Versi", "Dipakai untuk"], [[r["perangkat"], r["versi"] or "—", r["dipakai_untuk"]] for r in inv.get("perangkat", [])])
    table(["Alat", "Uji", "Hasil"], [[c["alat"], c["uji"], "LULUS" if c["lulus"] else "GAGAL"] for c in inv.get("verifikasi", [])], 2, ST_COLOR)
    doc.add_heading(f"3  {SECTIONS[2]}", 1)
    table(["Metode", "Aturan keputusan"], [[m["metode"], m["aturan"]] for m in g["analisis"]["metode"]])
    table(["TC", "Parameter terkunci"], [[t, json.dumps(v, ensure_ascii=False)] for t, v in run["parameter_terkunci"].items()] or [["—", "—"]])
    if run["deviasi"]:
        table(["TC", "Parameter", "Terkunci", "Dipakai", "Alasan"], [[x["tc"], x["parameter"], x["terkunci"], x["dipakai"], x["alasan"]]
                                                                    for x in run["deviasi"]])
    doc.add_heading(f"4  {SECTIONS[3]}", 1)
    table(["Versi", "LULUS", "GAGAL", "TERBLOKIR", "TDD", "Total", "Tingkat lulus"],
          [[f"v{v}", *[x["rekap"][s] for s in ("LULUS", "GAGAL", "TERBLOKIR", "TIDAK DAPAT DIUJI")], x["total"],
            f"{x['tingkat_lulus'] * 100:.1f}%"] for v, x in k["per_versi"].items()])
    for v, x in run["versi"].items():
        doc.add_paragraph(f"Versi {v}", style="Heading 3")
        table(["TC", "Judul", "Aktual", "Status"], [[t["id"], t["judul"], t["actual"], t["status"]] for t in x["tc"]], 3, ST_COLOR)
    doc.add_heading(f"5  {SECTIONS[4]}", 1)
    if not k["temuan"]:
        doc.add_paragraph("Tidak ada temuan ketidaksesuaian pada TC yang dijalankan.")
    for f in k["temuan"]:
        table(["Field", "Isi"], [["ID & judul", f"{f['id']} · {f['judul']} (v{f['versi']})"], ["Test case", f["test_case"]], ["Bukti", f["bukti"]],
                                 ["Dampak", f["dampak"]], ["Acuan", f["acuan"]], ["Keparahan", f["keparahan"]], ["Rekomendasi", f.get("rekomendasi") or "—"]])
        demo = f.get("demonstrasi_dampak") or {}
        if "berhasil" in demo:
            doc.add_paragraph(f"Bukti dampak: P2 “{demo['p2_diketahui']}” → P1 dipulihkan “{demo['p1_dipulihkan']}”.")
    doc.add_heading(f"6  {SECTIONS[5]}", 1)
    doc.add_paragraph(f"Verifikasi manifest: {g['verifikasi']['kesimpulan']}.")
    table(["Berkas", "SHA-256"], [[b["path"], b["sha256"]] for b in g["manifest"]["berkas"]])
    table(["Run ID", "Mulai", "Alasan"], [[x["run_id"], x["mulai"], x["alasan"]] for x in g["riwayat"]])
    doc.add_heading(f"7  {SECTIONS[6]}", 1)
    for v, x in k["per_versi"].items():
        doc.add_paragraph(f"Versi {v} — keseluruhan: {x['keseluruhan']}", style="Heading 3")
        table(["Sasaran", "Uraian", "Kategori", "Alasan"], [[s, o["nama"], o["kategori"], o["alasan"]] for s, o in sorted(x["sasaran"].items())], 2, CAT_COLOR)
    for s in k["pernyataan"]:
        doc.add_paragraph(s, style="List Bullet")
    table(["Peran", "Nama", "Tanda tangan", "Tanggal"], [["Penguji", penguji, "", ""], ["Pemeriksa", "", "", ""]])
    doc.add_paragraph("© Cryptan.ID — made by Dea Saka Kurnia Putra")
    doc.save(path)
    return Path(path)


def report(product_id, run_id, out_root=None, pdf=True, docx=True, penguji=None) -> dict:
    """Bangun laporan run, tambahkan ke manifest, lalu terbitkan salinan di outputs/<produk>/uk3/."""
    g = gather(product_id, run_id, out_root)
    d = g["dir"]
    out = {}
    hj = d / "uk3_hasil_uji.json"
    hj.write_text(json.dumps(hasil_uji_json(g), indent=1, ensure_ascii=False, default=str), encoding="utf-8")
    out["json"] = hj
    if pdf:
        out["pdf"] = build_pdf(g, d / "Laporan_Hasil_Pengujian.pdf", penguji)
    if docx:
        out["docx"] = build_docx(g, d / "Laporan_Hasil_Pengujian.docx", penguji)
    append(d, list(out.values()))
    pub = uk3_dir(product_id, out_root)
    for f in [hj, d / "lembar_hasil.csv", d / "evidence_manifest.json", *d.glob("log_*.txt"), *(v for k_, v in out.items() if k_ != "json")]:
        shutil.copy2(f, pub / f.name)
    if (pub / "evidence").exists():
        shutil.rmtree(pub / "evidence")
    shutil.copytree(d / "evidence", pub / "evidence")
    (pub / "TERBIT.txt").write_text(f"Salinan terbitan dari run {run_id} ({rel(d)}).\nVerifikasi: python -m uk3_pengujian "
                                    f"verify-evidence --product {product_id} --run {run_id}\n", encoding="utf-8")
    return {k_: rel(v) for k_, v in out.items()}
