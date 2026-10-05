"""std_report — katalog standar, plugin uji (kasus positif & negatif), skema results.json, laporan PDF/HTML/XLSX
dan konsistensi KPI laporan dengan results.json."""
import copy
import json

import pytest

from conftest import ROOT
from core.adapters import backends_for
from core.adapters.base import kind_of
from std_report import catalog, runner
from std_report.report.builder import build_pdf
from std_report.report.export import build_html, build_xlsx
from std_report.report.model import Model
from std_report.tests import DICATAT, GAGAL, INKON, LULUS, Ctx, registry
from std_report.tests import s as s_mod

PLUGINS = {p.id: p for p in registry()}
REQUIRED_DOCS = ["FIPS 180-4", "FIPS 186-5", "FIPS 197", "FIPS 198-1", "FIPS 202", "FIPS 203", "FIPS 204", "FIPS 205",
                 "SP 800-38A", "SP 800-38D", "SP 800-56A", "SP 800-90A", "SP 800-108", "SP 800-132", "SP 800-185",
                 "ISO/IEC 18033-3", "ISO/IEC 10118-3", "ISO/IEC 9797-1", "ISO/IEC 14888-3", "ISO/IEC 18032"]


@pytest.fixture(scope="module")
def cat():
    return catalog.build()


@pytest.fixture(scope="module")
def rows(cat):
    return {r["id"]: r for r in cat["rows"]}


def ctx(rows, rid, backends=None, row=None):
    r = row or rows[rid]
    return Ctx(r, kind_of(r), backends if backends is not None else backends_for(r), "ringan", runner.SEED)


def run(pid, c):
    return PLUGINS[pid].execute(c)


class Wrap:
    """Adapter tiruan: meneruskan semua atribut ke adapter asli kecuali yang ditimpa (untuk kasus negatif)."""

    def __init__(self, a, **over):
        self._a = a
        self.__dict__.update(over)

    def __getattr__(self, n):
        return getattr(self._a, n)


def _be(rows, rid, name):
    return next(b for b in backends_for(rows[rid]) if b.backend == name)


# --------------------------------------------------------------------------- katalog
def test_catalog_three_sources_and_source_doc(cat):
    assert cat["rows"], "katalog kosong"
    assert all(cat["counts"][b] > 0 for b in ("FIPS", "NIST-SP", "ISO-IEC"))
    assert all(r["source_doc"] and r["source_body"] for r in cat["rows"])
    assert cat["missing_docs"] == []
    for d in REQUIRED_DOCS:
        assert d in cat["doc_coverage"], d


def test_catalog_filters(cat):
    iso = catalog.build(["ISO-IEC"], ["hash"])
    assert iso["rows"] and all("ISO-IEC" in r["source_body"] and r["primitive"] == "hash" for r in iso["rows"])
    assert len(iso["rows"]) < len(cat["rows"])


# --------------------------------------------------------------------------- plugin: kasus positif
def test_plugins_positive_aes_cbc(rows):
    c = ctx(rows, "AES-128-CBC")
    got = {pid: run(pid, c)["status"] for pid in PLUGINS}
    expect = {"K-01": LULUS, "K-02": LULUS, "K-03": LULUS, "K-04": LULUS, "S-01": LULUS, "S-02": LULUS, "S-03": LULUS,
              "S-04": LULUS, "I-02": DICATAT, "I-03": LULUS}
    for pid, st in expect.items():
        assert got[pid] == st, (pid, got[pid])
    assert got["I-01"] in (LULUS, INKON)          # indikatif: bergantung pada derau waktu mesin


def test_k01_is_kat_when_vectors_exist(rows):
    r = run("K-01", ctx(rows, "SHA-256"))
    assert r["status"] == LULUS and r["label"] == "KAT"


def test_tdd_without_backend(rows):
    r = run("K-02", ctx(rows, "AES-128-CBC", backends=[]))
    assert r["status"] == "TIDAK_DAPAT_DIUJI" and r["reason"]


# --------------------------------------------------------------------------- plugin: kasus negatif
def test_k01_negative_wrong_digest(rows):
    bad = Wrap(_be(rows, "SHA-256", "stdlib"), digest=lambda m, o=None: bytes(32))
    assert run("K-01", ctx(rows, "SHA-256", [bad]))["status"] == GAGAL


def test_k02_negative_broken_decrypt(rows):
    bad = Wrap(_be(rows, "AES-128-CBC", "pyca"), decrypt=lambda k, n, d, aad=b"": d)
    assert run("K-02", ctx(rows, "AES-128-CBC", [bad]))["status"] == GAGAL


def test_k03_negative_backends_disagree(rows):
    bad = Wrap(_be(rows, "AES-128-CBC", "pyca"), encrypt=lambda k, n, d, aad=b"": bytes(len(d)))
    assert run("K-03", ctx(rows, "AES-128-CBC", [bad, _be(rows, "AES-128-CBC", "pycrypto")]))["status"] == GAGAL


def test_k04_negative_accepts_bad_key(rows):
    bad = Wrap(_be(rows, "AES-128-CBC", "pyca"), encrypt=lambda k, n, d, aad=b"": d)
    assert run("K-04", ctx(rows, "AES-128-CBC", [bad]))["status"] == GAGAL


def test_s01_negative_identity_cipher(rows):
    bad = Wrap(_be(rows, "AES-128-ECB", "pyca"), encrypt=lambda k, n, d, aad=b"": d)
    assert run("S-01", ctx(rows, "AES-128-ECB", [bad]))["status"] != LULUS


def test_s02_negative_constant_keystream(rows):
    bad = Wrap(_be(rows, "AES-128-CTR", "pyca"), encrypt=lambda k, n, d, aad=b"": bytes(len(d)))
    assert run("S-02", ctx(rows, "AES-128-CTR", [bad]))["status"] != LULUS


def test_s03_negative_wrong_reference(rows, monkeypatch):
    monkeypatch.setattr(s_mod, "AES_REF", {**s_mod.AES_REF, "NL": 100})
    assert run("S-03", ctx(rows, "AES-128-ECB"))["status"] == GAGAL


def test_s04_negative_weak_parameters(rows):
    r = copy.deepcopy(rows["AES-128-ECB"])
    r["security_strength_bits"] = 80
    assert run("S-04", ctx(rows, None, row=r))["status"] == GAGAL
    r["security_strength_bits"], r["status_nist"] = 128, "legacy_use"
    assert run("S-04", ctx(rows, None, row=r))["status"] == GAGAL


def test_i01_negative_secret_dependent_timing(rows):
    c = ctx(rows, "HMAC-SHA-256-K256")
    fixed = c.rand(64, "fixed")
    real = _be(rows, "HMAC-SHA-256-K256", "stdlib")

    def slow_mac(key, msg):
        if msg == fixed:                              # cabang bergantung data → beda waktu yang jelas
            sum(range(20000))
        return real.mac(key, msg)
    c.backends = [Wrap(real, mac=slow_mac)]
    assert run("I-01", c)["status"] == INKON


def test_i02_is_recorded_only(rows):
    r = run("I-02", ctx(rows, "SHA-256"))
    assert r["status"] == DICATAT and "ops/s" in r["actual"]


def test_i03_negative_crash(rows):
    def crash(*a, **k):
        raise SystemExit(3)
    bad = Wrap(_be(rows, "AES-128-CBC", "pyca"), decrypt=crash)
    assert run("I-03", ctx(rows, "AES-128-CBC", [bad]))["status"] == GAGAL


# --------------------------------------------------------------------------- results.json + laporan
@pytest.fixture(scope="module")
def small(tmp_path_factory):
    out = tmp_path_factory.mktemp("std_report")
    res = runner.run("ringan", only=["AES-128-ECB", "SHA-256", "HKDF", "SLH-DSA-SHA2-128S"], out_dir=out)
    return res, out


def test_results_schema(small):
    res, out = small
    assert res["schema_errors"] == []
    import jsonschema
    schema = json.loads(runner.SCHEMA.read_text())
    jsonschema.validate(json.loads((out / "results.json").read_text()), schema)
    assert all(t["test_id"] == f"{a['id']}/{t['plugin']}" for a in res["algorithms"] for t in a["tests"])
    assert all(len(a["tests"]) == len(PLUGINS) for a in res["algorithms"])   # tidak ada sel matriks kosong


def test_results_no_absolute_paths(small):
    _, out = small
    txt = (out / "results.json").read_text()
    assert str(ROOT) not in txt and "/Users/" not in txt


def test_kpi_consistency(small):
    res, _ = small
    m = Model(res)
    assert m.kpi() == res["summary"]["kpi"]
    assert m.status_counts() == res["summary"]["status_counts"]


def test_pdf_build(small, tmp_path):
    from pypdf import PdfReader
    res, _ = small
    p = build_pdf(res, out=tmp_path / "lap.pdf", penyusun="Penguji Pytest")
    rd = PdfReader(str(p))
    assert len(rd.pages) > 5
    cover = rd.pages[0].extract_text()
    assert "Laporan Daftar Algoritme" in cover and "Penguji Pytest" in cover and "TERBATAS" in cover
    assert rd.metadata.title.startswith("Laporan Daftar Algoritme") and rd.metadata.author == "Penguji Pytest"
    assert len(rd.outline) >= 9                                     # bookmark per bab + lampiran
    text = "\n".join(pg.extract_text() for pg in rd.pages)
    sc, k = res["summary"]["status_counts"], res["summary"]["kpi"]
    assert f"LULUS {sc['LULUS']}" in text and f"GAGAL {sc['GAGAL']}" in text     # Bab 1 = hitungan results.json
    assert f"{k['terdaftar']} kombinasi" in text
    assert f"Halaman 1 dari {len(rd.pages)}" not in cover                       # sampul tanpa footer
    assert f"Halaman 2 dari {len(rd.pages)}" in rd.pages[1].extract_text()
    ring = build_pdf(res, out=tmp_path / "ringkas.pdf", ringkas=True, penyusun="Penguji Pytest")
    assert len(PdfReader(str(ring)).pages) < len(rd.pages)


def test_html_xlsx(small, tmp_path):
    from openpyxl import load_workbook
    res, _ = small
    h = build_html(res, out=tmp_path / "lap.html", penyusun="Penguji Pytest")
    assert "TERBATAS" in h.read_text() and "Penguji Pytest" in h.read_text()
    x = build_xlsx(res, out=tmp_path / "lap.xlsx", penyusun="Penguji Pytest")
    wb = load_workbook(x)
    assert {"Ringkasan", "Daftar algoritme", "Matriks", "Hasil uji", "Vektor", "PERLU_VERIFIKASI"} <= set(wb.sheetnames)
    assert wb["Daftar algoritme"].max_row == len(res["algorithms"]) + 1


def test_report_filters(small, tmp_path):
    res, _ = small
    m = Model(res, sources="iso-iec", only_tested=True)
    assert all("ISO-IEC" in a["source_body"] and a["overall"] != "TIDAK_DAPAT_DIUJI" for a in m.algs)
    assert m.filename("pdf").name.endswith("_filter.pdf")
