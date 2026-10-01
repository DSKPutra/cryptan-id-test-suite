"""Modul algo_catalog — Daftar Algoritma yang Diuji (UK-1 KUK 1.1)."""
import io
import json
import shutil
import urllib.error
from pathlib import Path

import pytest
import yaml

from algo_catalog import catalog, extract, link, output, picker, scraper
from algo_catalog.normalize import Matcher, best, normalize
from algo_catalog.schema import PRIMITIVES, Entry, SchemaError
from algo_catalog.selection import Selection, warnings_for
from conftest import ROOT

SAMPLES = ROOT / "samples" / "input"


@pytest.fixture(scope="module")
def E():
    return catalog.load(use_cache=False)          # seed saja → deterministik


@pytest.fixture(scope="module")
def M(E):
    return Matcher(E)


def combo(E, cid):
    return catalog.combo_index(E)[cid.upper()]


# ---------------------------------------------------------------- schema & seed
def test_seed_valid_unique_and_covers_all_primitives(E):
    ids = [e.combo(k) for e in E for k in e.keys()]
    assert len(ids) == len(set(i.upper() for i in ids)), "ID kanonik kombinasi duplikat"
    assert {e.primitive for e in E} == set(PRIMITIVES)


@pytest.mark.parametrize("fam,variants", [
    ("AES", {"ECB", "CBC", "CFB1", "CFB8", "CFB128", "OFB", "CTR", "CMAC", "CCM", "GCM", "GMAC", "XTS", "KW", "KWP", "FF1", "FF3-1"}),
    ("SLH-DSA", {f"{h}-{s}{v}" for h in ("SHA2", "SHAKE") for s in ("128", "192", "256") for v in "sf"}),
    ("ML-DSA", {"ML-DSA-44", "ML-DSA-65", "ML-DSA-87"}),
    ("ML-KEM", {"ML-KEM-512", "ML-KEM-768", "ML-KEM-1024"}),
    ("DRBG", {"Hash_DRBG", "HMAC_DRBG", "CTR_DRBG"}),
])
def test_seed_minimum_contents(E, fam, variants):
    assert variants <= {e.variant for e in E if e.family == fam}


def test_seed_required_algorithms_present(E):
    have = {e.family for e in E}
    for f in ["TDEA", "Camellia", "SEED", "ARIA", "SM4", "HIGHT", "PRESENT", "CLEFIA", "LEA", "ChaCha", "Salsa20", "SNOW",
              "MUGI", "Rabbit", "Trivium", "Enocoro", "Grain", "ZUC", "RC4", "SHA-1", "SHA-2", "SHA-3", "Ascon", "SM3",
              "Streebog", "RIPEMD", "Whirlpool", "BLAKE2", "HMAC", "KMAC", "Poly1305", "RSA", "ECIES", "ECDH", "XECDH",
              "FFDH", "ElGamal", "ECDSA", "EdDSA", "DSA", "LMS", "XMSS", "SM2", "KDF"]:
        assert f in have, f


def test_schema_rejects_bad_entries():
    with pytest.raises(SchemaError):
        Entry.from_dict({"id": "X", "primitive": "quantum", "family": "X", "variant": "X"})
    with pytest.raises(SchemaError):
        Entry.from_dict({"id": "X", "primitive": "hash", "family": "X", "variant": "X", "status_nist": "ok"})
    with pytest.raises(SchemaError):
        Entry.from_dict({"id": "X", "primitive": "hash", "family": "X", "variant": "X", "foo": 1})


# ---------------------------------------------------------------- strength & status
@pytest.mark.parametrize("cid,bits,cat", [
    ("RSA-OAEP-2048", 112, None), ("RSA-OAEP-3072", 128, None), ("RSA-PSS-2048", 112, None),
    ("ECDSA-P384", 192, None), ("ECDH-P384", 192, None), ("ECDSA-P256", 128, None), ("ECDH-P224", 112, None),
    ("ML-KEM-768", 192, 3), ("ML-KEM-512", 128, 1), ("ML-DSA-44", 128, 2), ("SLH-DSA-SHAKE-256F", 256, 5),
    ("AES-256-GCM", 256, None), ("XTS-AES-512", 256, None), ("SHA3-256", 128, None), ("SHA-512/224", 112, None),
    ("TDEA-3KEY", 112, None), ("FFDHE2048", 112, None), ("ED448", 224, None),
])
def test_security_strength(E, cid, bits, cat):
    e, k = combo(E, cid)
    assert e.strength(k) == bits
    assert e.category(k) == cat


@pytest.mark.parametrize("cid,status", [
    ("RSA-OAEP-1024", "disallowed"), ("RSA-OAEP-2048", "acceptable"), ("TDEA-3KEY", "disallowed"),
    ("RC4-128", "disallowed"), ("SHA-1", "disallowed"), ("DSA-2048", "legacy_use"), ("RSAES-PKCS1-V1_5-2048", "disallowed"),
    ("ECDSA-SECP256K1", "not_nist"), ("CAMELLIA-128", "not_nist"), ("ML-KEM-1024", "acceptable"), ("AES-128-FF3-1", "PERLU_VERIFIKASI"),
])
def test_status_nist(E, cid, status):
    e, k = combo(E, cid)
    assert e.status(k) == status


# ---------------------------------------------------------------- normalisasi alias
@pytest.mark.parametrize("text,key,expected", [
    ("aes256gcm", None, ["AES-256-GCM"]),
    ("AES/GCM/NoPadding", 256, ["AES-256-GCM"]),
    ("EVP_aes_256_gcm", None, ["AES-256-GCM"]),
    ("id-aes256-GCM", None, ["AES-256-GCM"]),
    ("AES_128_CBC", None, ["AES-128-CBC"]),
    ("hashlib.sha3_256(data)", None, ["SHA3-256"]),
    ("SHA-512/256", None, ["SHA-512/256"]),
    ("HMAC-SHA3-256", 256, ["HMAC-SHA3-256-K256"]),
    ("XTS-AES-256", None, ["XTS-AES-256"]),
    ("Kyber768", None, ["ML-KEM-768"]),
    ("RSA-OAEP 2048-bit and 3072-bit", None, ["RSA-OAEP-2048", "RSA-OAEP-3072"]),
    ("Cipher.getInstance(\"RSA/ECB/OAEPWithSHA-256AndMGF1Padding\"); g.initialize(4096)", None, ["RSA-OAEP-4096"]),
    ("SPHINCS+-SHAKE-128f", None, ["SLH-DSA-SHAKE-128F"]),
    ("CTR_DRBG-AES-256", None, ["CTR-DRBG-AES-256"]),
])
def test_alias_normalization(M, text, key, expected):
    assert sorted(normalize(M, text, key)) == sorted(expected)


def test_no_false_positives(M):
    for t in ["a random seed value", "the DSA in ECDSA", "AES-128-CTR, AES-256-GCM", "we use present tense"]:
        ids = normalize(M, t)
        assert "SEED-128" not in ids and "DSA-2048" not in ids and "AES-256-CTR" not in ids and "PRESENT-80" not in ids


def test_containment_prefers_longest_match(M):
    assert normalize(M, "HMAC-SHA-256 key 128") == ["HMAC-SHA-256-K128"]
    assert normalize(M, "AES-256-GCM-SIV") == ["AES-256-GCM-SIV"]


def test_ambiguous_curve_uses_context_hint(M):
    d = {x["id"]: x["confidence"] for x in best(M.find("key = ec.generate_private_key(ec.SECP384R1()); key.sign(m, ECDSA)"))}
    assert d["ECDSA-P384"] > d["ECDH-P384"] and d["ECDSA-P384"] >= 0.6 > d["ECDH-P384"]


# ---------------------------------------------------------------- ekstraksi file
def _ids(M, path):
    return {d["id"] for d in best(M.find(extract.file_to_text(path))) if d["confidence"] >= 0.6}


def test_extract_pdf_sample(M):
    ids = _ids(M, SAMPLES / "datasheet_securelib.pdf")
    assert {"AES-128-CTR", "AES-256-GCM", "ML-KEM-768", "RSA-OAEP-2048", "RSA-OAEP-3072", "ECDSA-P256", "ML-DSA-65", "TDEA-3KEY"} <= ids


def test_extract_html_sample_ignores_script(M):
    ids = _ids(M, SAMPLES / "vendor_page.html")
    assert {"CHACHA20-POLY1305", "XTS-AES-256", "SLH-DSA-SHA2-128S", "X25519", "FFDHE3072", "ASCON-AEAD128"} <= ids
    assert "AES-128-ECB" not in ids


@pytest.mark.parametrize("name,expected", [
    ("securelib_crypto.c", {"AES-256-GCM", "TDEA-3KEY", "SHA3-256", "ECDSA-P256"}),
    ("CryptoService.java", {"AES-256-GCM", "RSA-OAEP-3072"}),
    ("app_crypto.py", {"SHA3-256", "SHA-1", "ECDSA-P384", "ED25519"}),
])
def test_extract_source_code(M, name, expected):
    assert expected <= _ids(M, SAMPLES / name)


def test_extract_office_and_structured(tmp_path, M):
    import docx
    from openpyxl import Workbook
    d = docx.Document(); d.add_paragraph("Algoritma: SM4 dan SM3"); t = d.add_table(rows=1, cols=1); t.rows[0].cells[0].text = "Ed448"
    d.save(tmp_path / "spec.docx")
    wb = Workbook(); wb.active.append(["Hash", "BLAKE2s"]); wb.save(tmp_path / "alg.xlsx")
    (tmp_path / "a.json").write_text(json.dumps({"ciphers": ["ChaCha20", {"mac": "Poly1305"}]}))
    (tmp_path / "a.yaml").write_text("kem: [ML-KEM-1024]\n")
    (tmp_path / "a.csv").write_text("alg,key\nCamellia-192,192\n")
    assert {"SM4-128", "SM3", "ED448"} <= _ids(M, tmp_path / "spec.docx")
    assert "BLAKE2S-256" in _ids(M, tmp_path / "alg.xlsx")
    assert {"CHACHA20-256", "POLY1305"} <= _ids(M, tmp_path / "a.json")
    assert "ML-KEM-1024" in _ids(M, tmp_path / "a.yaml")
    assert "CAMELLIA-192" in _ids(M, tmp_path / "a.csv")
    with pytest.raises(extract.ExtractError):
        extract.bytes_to_text("x.exe", b"MZ")


# ---------------------------------------------------------------- dropdown / pick
@pytest.mark.parametrize("spec,n", [
    ("AES:GCM:128,256", 2), ("AES:GCM", 3), ("ECDSA:P-384", 1), ("ML-KEM:768", 1), ("AES-256-GCM", 1),
    ("SLH-DSA:*", 12), ("RSA:OAEP", 6), ("SHA-3:SHAKE128", 1),
])
def test_pick_parsing(E, spec, n):
    assert len(picker.parse_pick(spec, E)) == n


def test_pick_all_variants_counts_combinations(E):
    aes = [e for e in E if e.family == "AES"]
    assert len(picker.parse_pick("AES:*", E)) == sum(len(e.keys()) for e in aes)


@pytest.mark.parametrize("spec", ["AES:GCM:512", "FOO:BAR", "AES:ZZZ", "NOT-AN-ID"])
def test_pick_errors(E, spec):
    with pytest.raises(picker.PickError):
        picker.parse_pick(spec, E)


def test_combination_count(E):
    c = catalog.count_combinations(E)
    assert c["total"] == sum(len(e.keys()) for e in E) == sum(v for k, v in c.items() if k != "total")


# ---------------------------------------------------------------- seleksi: gabung, dedup, konfirmasi, peringatan
def test_merge_dedup_across_inputs(E, M):
    sel = Selection(E)
    sel.add_pick(picker.parse_pick("AES:GCM:256", E), "AES:GCM:256")
    sel.add_detections(best(M.find("EVP_aes_256_gcm and aes256gcm")), "file", "x.c")
    sel.add_detections(best(M.find("AES-256-GCM")), "link", "https://example.org")
    assert list(sel.items) == ["AES-256-GCM"]
    it = sel.items["AES-256-GCM"]
    assert {i["type"] for i in it["inputs"]} == {"dropdown", "file", "link"} and it["confirmed"]


def test_detections_require_confirmation(E, M):
    sel = Selection(E)
    sel.add_detections(best(M.find("SHA3-256 and Mac.getInstance(\"HmacSHA384\")")), "file", "a.java")
    assert sel.final() == [] and len(sel.pending()) == 3
    sel.confirm(["SHA3-256"])
    assert [r["id"] for r in sel.final()] == ["SHA3-256"]
    sel2 = Selection(E)
    sel2.add_detections(best(M.find("SHA3-256 and HmacSHA384")), "file", "a.java", auto_confirm_at=0.6)
    assert [r["id"] for r in sel2.final()] == ["SHA3-256"]


def test_edit_remove_and_change_key(E):
    sel = Selection(E)
    sel.add_pick(picker.parse_pick("RSA:OAEP:2048", E), "pick")
    nid = sel.change_key("RSA-OAEP-2048", 3072)
    assert nid == "RSA-OAEP-3072" and list(sel.items) == ["RSA-OAEP-3072"]
    with pytest.raises(ValueError):
        sel.change_key("RSA-OAEP-3072", 999)
    sel.remove(["RSA-OAEP-3072"])
    assert not sel.items


@pytest.mark.parametrize("cid,codes", [
    ("PRESENT-80", {"STRENGTH_LT_112"}), ("TDEA-3KEY", {"STATUS_DISALLOWED"}),
    ("SHA-1", {"STATUS_DISALLOWED", "STRENGTH_LT_112"}), ("DSA-2048", {"STATUS_LEGACY_USE", "QUANTUM_VULNERABLE"}),
    ("RSA-OAEP-2048", {"QUANTUM_VULNERABLE"}), ("ML-KEM-768", set()), ("AES-256-GCM", set()),
])
def test_warnings(E, cid, codes):
    e, k = combo(E, cid)
    got = {w["code"] for w in warnings_for(e.combo_row(k))}
    assert codes <= got and (codes or not got - {"PERLU_VERIFIKASI"})
    if "QUANTUM_VULNERABLE" in codes:
        assert any("8547" in w["message"] for w in warnings_for(e.combo_row(k)))


# ---------------------------------------------------------------- link & scraper (tanpa jaringan)
class _Resp(io.BytesIO):
    def __init__(self, data, ctype):
        super().__init__(data)
        self.headers = {"Content-Type": ctype}

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False


def _opener(routes):
    def op(req):
        url = req.full_url
        if url not in routes:
            raise urllib.error.URLError("offline")
        data, ctype = routes[url]
        return _Resp(data, ctype)
    return op


def test_link_fetch_html_and_pdf(monkeypatch):
    link._robots.clear()
    pdf = (SAMPLES / "datasheet_securelib.pdf").read_bytes()
    op = _opener({"https://ex.test/robots.txt": (b"User-agent: *\nDisallow: /private\n", "text/plain"),
                  "https://ex.test/page": (b"<html><body><p>Uses SHA3-512</p></body></html>", "text/html"),
                  "https://ex.test/ds.pdf": (pdf, "application/pdf")})
    page = link.fetch("https://ex.test/page", delay=0, opener=op)
    assert page["kind"] == "html" and "SHA3-512" in page["text"]
    assert link.fetch("https://ex.test/ds.pdf", delay=0, opener=op)["kind"] == "pdf"
    with pytest.raises(link.FetchError, match="robots"):
        link.fetch("https://ex.test/private/x", delay=0, opener=op)
    with pytest.raises(link.FetchError):
        link.fetch("ftp://ex.test/x", delay=0, opener=op)


def test_scraper_offline_keeps_seed_and_old_cache(tmp_path):
    cache = tmp_path / "cache.json"
    old = {"entries": [{"id": "ACVP-OLD", "primitive": "hash", "family": "Old", "variant": "Old", "status_nist": "PERLU_VERIFIKASI",
                        "source": {"type": "scrape", "source_id": "acvp-index"}}],
           "updates": [{"id": "AES-GCM", "url": "https://old", "source_id": "cavp"}], "sources": []}
    cache.write_text(json.dumps(old))

    def offline(url, delay=0, **k):
        raise link.FetchError("offline")
    res = scraper.scrape(refresh=True, cache_path=cache, fetcher=offline, raw_fetcher=offline, delay=0, log=lambda *a: None)
    assert all(s["status"] == "error" for s in res["sources"])
    assert [e["id"] for e in res["entries"]] == ["ACVP-OLD"] and res["updates"][0]["url"] == "https://old"
    merged = catalog.load(cache=cache)
    assert len(merged) == len(catalog.load(use_cache=False)) + 1


def test_scraper_acvp_index_adds_and_matches(tmp_path, E):
    html = "".join(f'<li><a href="https://pages.nist.gov/ACVP/draft-x.txt">{n}</a></li>'
                   for n in ["AES-GCM", "SHA3-256 2.0", "ML-DSA sigGen", "TupleHash-128", "ECDSA mode: keyGen"])
    cache = tmp_path / "c.json"
    res = scraper.scrape(refresh=True, cache_path=cache, raw_fetcher=lambda u, d: html,
                         fetcher=lambda u, delay=0: {"text": "AES-256-CBC", "kind": "html"}, delay=0, log=lambda *a: None)
    assert [e["variant"] for e in res["entries"]] == ["TupleHash-128"]
    assert res["entries"][0]["status_nist"] == "PERLU_VERIFIKASI" and res["entries"][0]["primitive"] == "hash"
    upd = {u["id"] for u in res["updates"]}
    assert {"AES-GCM", "SHA3-256", "ML-DSA-44", "ECDSA-P256"} <= upd
    assert res["generated"] and all(s["retrieved"] for s in res["sources"])


# ---------------------------------------------------------------- keluaran & tautan UK-1
def test_outputs_and_profile_link(tmp_path, E, monkeypatch):
    sel = Selection(E)
    for spec in ["AES:CTR:128", "PRESENT:ECB:80", "ECDSA:P-256", "ML-KEM:768"]:
        sel.add_pick(picker.parse_pick(spec, E), spec)
    prof = tmp_path / "config" / "product_profile.yaml"
    prof.parent.mkdir()
    shutil.copy(ROOT / "config" / "product_profile.yaml", prof)
    monkeypatch.setattr(output, "PROFILE", prof)
    res = output.write_all(sel.final(), tmp_path / "out", link=True)
    for k in ("yaml", "md", "csv", "xlsx"):
        assert res["files"][k].exists()
    doc = yaml.safe_load(res["files"]["yaml"].read_text())
    assert doc["summary"]["total_combinations"] == 4
    assert {i["id"]: i["uk1_profile"] for i in doc["items"]}["AES-128-CTR"] == "config/profiles/aes128.yaml"
    assert any(w["code"] == "STRENGTH_LT_112" for w in doc["warnings"])
    md = res["files"]["md"].read_text()
    for col in output.COLUMNS:
        assert col in md
    assert "algorithms_under_test: ../out/algorithms_under_test.yaml" in prof.read_text()
    real = (ROOT / "config" / "product_profile.yaml").read_text()
    assert "pytest-of" not in real and "/private/" not in real       # profil asli tidak tersentuh


def test_uk1_profile_reads_linked_list():
    from uk1_metode import profile
    from conftest import profile_of
    s = profile.summarize(profile_of("AES-128"))["algorithms_under_test"]
    if not s["linked"]:
        pytest.skip("outputs/algo_catalog belum dibangkitkan")
    assert s["catalog_id"] == "AES-128-CTR" and s["in_list"] and not s["file"].startswith("/")


def test_cli_select_pick(tmp_path):
    from algo_catalog.__main__ import main
    assert main(["select", "--pick", "AES:GCM:128,256", "ECDSA:P-384", "--out", str(tmp_path), "--no-link-profile", "--no-cache"]) == 0
    doc = yaml.safe_load((tmp_path / output.YAML_NAME).read_text())
    assert sorted(i["id"] for i in doc["items"]) == ["AES-128-GCM", "AES-256-GCM", "ECDSA-P384"]


def test_streamlit_app_smoke():
    pytest.importorskip("streamlit")
    from streamlit.testing.v1 import AppTest
    at = AppTest.from_file(str(ROOT / "algo_catalog" / "app.py"), default_timeout=60).run()
    assert not at.exception and len(at.tabs) == 3
