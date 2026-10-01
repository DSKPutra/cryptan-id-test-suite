"""KUK 2.5 — Kompilasi skenario.

Menggabungkan: dokumen skenario (Bab 1–9), katalog test case, matriks keterlacakan
(kebutuhan → metode UK-1 → skenario → test case → area/standar), dan tabel pemetaan
parameter → kemungkinan hasil (Memenuhi / Memenuhi dengan Catatan / Tidak Memenuhi /
Inkonklusif) beserta parameter pemicu & tindak lanjut. Model 4 tingkat: bab tata kelola.
"""
import csv
import json
from pathlib import Path

from uk1_metode.report import H, L, NOTE, P, T, _fmt, to_docx, to_markdown

from . import verifier as V
from .context import LAYERS, rel
from .param_space import TECH

SCHEMA = Path(__file__).parent / "schema" / "uk2_skenario.schema.json"
OUT_JSON, OUT_MD, OUT_DOCX = "uk2_skenario.json", "UK2_Dokumen_Skenario_Pengujian.md", "UK2_Dokumen_Skenario_Pengujian.docx"


# --------------------------------------------------------------------- keterlacakan
def traceability(ctx, scenarios, tcs) -> list:
    uk1_req = {}
    for r in ctx.uk1:
        alg = r.get("profile", {}).get("algorithm", {}).get("id")
        for t in r.get("traceability", []):
            uk1_req.setdefault(t["method_id"], []).append({"id": f"{alg}:{t['id']}", "text": t["requirement"]})
    sc_of = {rid: s["id"] for s in scenarios for rid in s["recipes"]}
    claims = {tag: c for c in ctx.claims for tag in c.get("tags", [])}
    rows = []
    for tc in tcs:
        reqs = []
        for m in tc["methods"]:
            reqs += [(q["id"], q["text"], m) for q in uk1_req.get(m, [])[:2]]
        for tag in tc["claims"]:
            if tag in claims:
                reqs.append((claims[tag]["id"], claims[tag]["claim"], tc["methods"][0] if tc["methods"] else "—"))
        if not reqs:
            reqs = [(f"OBJ:{tc['targets'][0]}", f"Kesesuaian {', '.join(tc['targets'][:3])} terhadap standar", tc["methods"][0] if tc["methods"] else "—")]
        for rid, text, m in reqs:
            rows.append({"requirement_id": rid, "requirement": text, "method": m, "scenario": sc_of.get(tc["id"], "—"),
                         "test_case": tc["id"], "area": tc.get("area") or "—", "standards": ", ".join(tc["standards"]),
                         "targets": ", ".join(tc["targets"][:6])})
    return rows


# --------------------------------------------------------------------- pemetaan parameter → hasil
GENERIC_OUTCOMES = [
    {"kind": "deterministic", "condition": "Keluaran berbeda dari expected value pada ≥ 1 vektor", "trigger": "target, vektor (kunci/nonce/pesan)",
     "outcome": "Tidak Memenuhi", "action": "Telusuri vektor pertama yang gagal; perbaiki implementasi; ulangi KAT/MCT"},
    {"kind": "deterministic", "condition": "Vektor 'invalid' (Wycheproof) diterima produk", "trigger": "encoding/parameter invalid",
     "outcome": "Tidak Memenuhi", "action": "Temuan kritis: perketat validasi masukan; ulangi seluruh vektor invalid"},
    {"kind": "deterministic", "condition": "Expected value bertanda PERLU_VERIFIKASI (vektor resmi belum tersedia)", "trigger": "ketersediaan vektor",
     "outcome": "Inkonklusif", "action": "Muat vektor ACVP/CAVP resmi, regenerasi expected, ulangi uji"},
    {"kind": "statistical", "condition": "Proporsi lolos < ambang SP 800-22 atau P-value_T < 0,0001", "trigger": "pola masukan, ukuran sampel",
     "outcome": "Tidak Memenuhi", "action": "Identifikasi dataset pemicu (low/high-density, avalanche); analisis difusi"},
    {"kind": "statistical", "condition": "Proporsi lolos tepat di tepi interval kepercayaan", "trigger": "ukuran sampel",
     "outcome": "Inkonklusif", "action": "Perbesar sampel (mis. 10×) lalu ulangi"},
    {"kind": "cryptanalytic", "condition": "Security strength 112 ≤ s < 128 bit", "trigger": "ukuran kunci/domain",
     "outcome": "Memenuhi dengan Catatan", "action": "Batasi penggunaan hingga 2030 (SP 800-131A / IR 8547); rencanakan migrasi"},
    {"kind": "cryptanalytic", "condition": "Serangan terbaik < klaim bit security / strength < 112 bit", "trigger": "ukuran kunci, jumlah ronde",
     "outcome": "Tidak Memenuhi", "action": "Nonaktifkan untuk layanan approved; ganti algoritma/parameter"},
    {"kind": "implementation", "condition": "TVLA |t| ≥ 4,5", "trigger": "kunci/nonce rahasia, jumlah pengukuran",
     "outcome": "Tidak Memenuhi", "action": "Konfirmasi eksploitabilitas (CPA/analisis timing); perbaiki ke constant-time"},
    {"kind": "implementation", "condition": "Hasil TVLA tidak konsisten antarsesi", "trigger": "jumlah trace/pengukuran, derau",
     "outcome": "Inkonklusif", "action": "Perbesar jumlah pengukuran; kendalikan derau (isolasi core, turbo off)"},
    {"kind": "implementation", "condition": "Crash/hang pada fuzzing", "trigger": "masukan malformed", "outcome": "Tidak Memenuhi",
     "action": "Reproduksi dengan sanitizer; perbaiki parser; ulangi fuzzing"},
    {"kind": "negative", "condition": "Masukan negatif diterima", "trigger": "nilai di luar spesifikasi", "outcome": "Tidak Memenuhi",
     "action": "Tambahkan validasi; ulangi uji negatif"},
    {"kind": "negative", "condition": "Masukan ditolak tetapi kode galat berbeda antar-kelas", "trigger": "jenis kegagalan",
     "outcome": "Memenuhi dengan Catatan", "action": "Seragamkan kode galat (mitigasi oracle) bila menyangkut data rahasia"},
]


def outcome_map(ctx, tcs) -> list:
    rows = []
    ids = {t["id"] for t in tcs}
    for name, ex in ctx.examples.items():                    # tabel materi bila contoh relevan dengan profil
        ex_ids = {t["id"] for t in ex.get("test_cases", [])}
        if ex_ids and ex_ids <= ids:
            rows += [{**o, "source": f"materi ({name})", "kind": "contoh"} for o in ex.get("outcome_map", [])]
    have = {t["expected_kind"] for t in tcs}
    rows += [{**{k: v for k, v in g.items() if k != "kind"}, "kind": g["kind"], "source": "aturan umum (kriteria keputusan materi)"}
             for g in GENERIC_OUTCOMES if g["kind"] in have]
    for o in ctx.objects:                                     # baris khusus dari peringatan algo_catalog
        st, s = o.get("status_nist"), o.get("security_strength_bits")
        if st in ("disallowed", "deprecated", "legacy_use"):
            rows.append({"condition": f"{o['id']} dipakai untuk layanan approved", "trigger": "status NIST " + st, "outcome": "Tidak Memenuhi",
                         "action": "Batasi ke dekripsi/verifikasi legacy; Memenuhi dengan Catatan bila hanya legacy use", "kind": "profil", "source": "algorithms_under_test"})
        elif isinstance(s, int) and s < 112:
            rows.append({"condition": f"{o['id']} strength {s} bit < 112", "trigger": "ukuran kunci", "outcome": "Tidak Memenuhi",
                         "action": "Ganti ke varian ≥ 128 bit", "kind": "profil", "source": "algorithms_under_test"})
        elif o.get("quantum_vulnerable") and isinstance(s, int) and s < 128:
            rows.append({"condition": f"{o['id']} ({s} bit) rentan kuantum", "trigger": "ukuran kunci/domain", "outcome": "Memenuhi dengan Catatan",
                         "action": "Diterima hingga 2030 (IR 8547); migrasi ke ML-KEM/ML-DSA", "kind": "profil", "source": "algorithms_under_test"})
    return rows


# --------------------------------------------------------------------- dokumen
def build_document(R: dict) -> list:
    pf, tcs = R["profile"], R["test_cases"]
    D = [H(1, f"Dokumen Skenario Pengujian — {pf['product']['name']}"),
         P(f"**{R['app']['name']}** · Modul **UK-2** · Unit {R['unit']['code']} — *{R['unit']['title']}* · {R['unit']['skkni']}"),
         P(f"Dibangkitkan {R['generated_at']} · model {pf['testing_model']} tingkat · masukan UK-1: `{R['inputs']['uk1']}`"),
         NOTE("Objek uji & data produk ILUSTRATIF. HASIL UJI LANGSUNG = dihitung kode; LITERATUR/ACUAN = dikutip dengan rujukan; "
              "PERLU_VERIFIKASI = belum dapat dihitung/dicocokkan, tidak dikarang."),
         H(2, "Ringkasan"),
         L([f"Objek uji: **{len(pf['objects'])}** kombinasi · kategori: {', '.join(pf['categories'])}",
            f"Metode UK-1 ditelaah: {R['method_review']['summary']['methods']} (relevan {R['method_review']['summary']['relevant']}) · sel kosong diisi templat: {R['method_review']['summary']['cells_empty']}",
            f"Skenario: **{len(R['scenarios'])}** · test case: **{len(tcs)}** (K {sum(t['layer'] == 'K' for t in tcs)} · S {sum(t['layer'] == 'S' for t in tcs)} · I {sum(t['layer'] == 'I' for t in tcs)})",
            f"Verifikasi terhadap spesifikasi desain: **{R['verification']['status']}** ({R['verification']['total_findings']} temuan)",
            f"Expected value: {R['expected']['summary']['COCOK_VEKTOR_RESMI']} cocok vektor resmi · {R['expected']['summary']['KRITERIA']} kriteria · {R['expected']['summary']['PERLU_VERIFIKASI']} PERLU_VERIFIKASI"]) ]
    # 1
    D += [H(2, "1. Profil Objek Uji & Ruang Parameter"),
          T(["Atribut", "Nilai"], [["Produk", f"{pf['product']['name']} {pf['product'].get('version', '')}"], ["Deskripsi", pf["product"].get("description", "")],
                                    ["Jenis produk", pf["product_type"]], ["Model tingkat", f"{pf['testing_model']} ({' → '.join(pf['tiers'])})"],
                                    ["SL ISO/IEC 19790", pf.get("security_level") or "tidak berlaku"], ["Fitur", ", ".join(pf["features"])]]),
          P("**Klaim keamanan:**"), T(["ID", "Klaim", "Tag"], [[c["id"], c["claim"], ", ".join(c.get("tags", []))] for c in pf["claims"]] or [["—", "—", "—"]]),
          P("**Objek uji (algorithms_under_test):**"),
          T(["ID", "Kategori", "Varian", "Kunci", "Strength", "Status NIST"], [[o["id"], o["category"], o["variant"], o.get("key_bits") or "—",
                                                                                o.get("security_strength_bits"), o.get("status_nist")] for o in pf["objects"]])]
    for cat, c in R["param_space"]["categories"].items():
        D += [H(3, f"1.{list(R['param_space']['categories']).index(cat) + 1} Ruang parameter — {cat}"),
              T(["ID", "Parameter", "Kelas", "Status", "Nilai uji (teknik)"],
                [[p["id"], p["name"], p["class"], p["status"], "; ".join(f"{x['value']} [{x['technique']}{'' if x['in_spec'] else ', NEG'}]" for x in p["values"])]
                 for p in c["params"]]),
              P(f"Reduksi pairwise (t = 2, IPOG) atas {len(c['pairwise']['parameters'])} parameter Variabel: "
                f"**{c['pairwise']['full_combinations']:,}** kombinasi penuh → **{c['pairwise']['pairwise_combinations']}** kombinasi pairwise "
                f"({c['pairwise']['reduction_pct']}% berkurang; semua pasangan tercakup: {c['pairwise']['all_pairs_covered']}).".replace(",", "."))]
    # 2
    mr = R["method_review"]
    D += [H(2, "2. Telaah Metode & Parameter (KUK 1.1, 1.2)"),
          P("Empat kebutuhan pengujian SKKNI: N1 jenis algoritma · N2 desain & teknik implementasi · N3 tren serangan · N4 best practice."),
          T(["Metode", "Sel (tingkat × lapis)", "N1", "N2", "N3", "N4", "Asal UK-1", "Putusan"],
            [[f"{m['method_id']} {m['name']}", f"{m['tier']} × {m['layer']}", *["✔" if m["needs"][k] else "✘" for k in m["needs"]],
              ", ".join(x for x in m["uk1_algorithms"] if x) + (" (STUB_UK1)" if m["stub"] else ""), m["verdict"]] for m in mr["methods"]]),
          P(f"Sel kosong ({len(mr['empty_cells'])}) diisi dari templat kategori materi: "
            + ", ".join(f"{e['category']}:{e['tier']}×{e['layer']}" for e in mr["empty_cells"][:30]) + ("…" if len(mr["empty_cells"]) > 30 else "")),
          P("Teknik penetapan nilai uji: " + "; ".join(f"{k} = {v}" for k, v in TECH.items()))]
    # 3
    D += [H(2, "3. Matriks Skenario Lapis × Tingkat (KUK 2.1)")]
    for g, m in R["design"]["matrix"].items():
        rows = []
        for t, cells in m.items():
            row = [t]
            for Lr in "KSI":
                sid = cells.get(Lr)
                sc = next((s for s in R["scenarios"] if s["id"] == sid), None)
                row.append(f"{sid}: {', '.join(sc['recipes']) or sc['status']}" if sc else "—")
            rows.append(row)
        D += [P(f"**{g}**"), T(["Tingkat", LAYERS["K"], LAYERS["S"], LAYERS["I"]], rows)]
    cx = R["design"]["cross"]
    D += [P("**Lintas-algoritma:** " + " → ".join(cx["chain"]) + f" · skenario: {', '.join(cx['recipes'])} · PQC relevan: {cx['pqc_relevant']}")]
    if R["design"].get("excluded"):
        D += [P("**Resep tidak berlaku (fitur/objek tidak dimiliki produk — aturan f):**"),
              T(["Resep", "Alasan"], [[e["id"], e["reason"]] for e in R["design"]["excluded"]])]
    if R["design"].get("iso_mapping"):
        D += [P("**Pemetaan area ISO/IEC 19790:**"),
              T(["Area", *pf["tiers"]], [[a, *[", ".join(v[t]) or "—" for t in pf["tiers"]]] for a, v in R["design"]["iso_mapping"].items()])]
    # 4
    ver = R["verification"]
    D += [H(2, "4. Laporan Verifikasi terhadap Spesifikasi Desain (KUK 2.2 — aspek kritis)"),
          P(f"Status keseluruhan: **{ver['status']}** · {ver['total_findings']} temuan"),
          T(["Aturan", "Deskripsi", "Diperiksa", "Status", "Temuan"], [[f"({r['rule']})", r["desc"], r["checked"], r["status"],
                                                                         "; ".join(f["detail"] for f in r["findings"]) or "—"] for r in ver["rules"]])]
    # 5
    D += [H(2, "5. Katalog Test Case (KUK 2.3)")]
    for Lr in "KSI":
        D += [H(3, f"5.{'KSI'.index(Lr) + 1} {LAYERS[Lr]}"),
              T(["ID", "Tingkat", "Area", "Parameter & nilai uji", "Prosedur", "Kriteria lulus", "Metode UK-1", "Prioritas", "Pelaksana"],
                [[t["id"], t["tier"] + (f" ({t['tier_note']})" if t.get("tier_note") else ""), t.get("area") or "—",
                  (t["material_reference"]["params"] + " · " if t.get("material_reference") else "")
                  + f"{len(t['parameters'])} nilai ({', '.join(sorted({p['name'] for p in t['parameters']})[:3])})" if t["parameters"] else
                  (t["material_reference"]["params"] if t.get("material_reference") else ", ".join(t["targets"][:4])),
                  " ".join(t["steps"]), t["pass"], ", ".join(t["methods"]), t["priority"], t["executor"]] for t in R["test_cases"] if t["layer"] == Lr])]
    D += [P("**Narasi empiris (contoh):**"), L([f"**{t['id']}** — {t['narrative']}" for t in R["test_cases"][:8]])]
    # 6
    ex = R["expected"]
    D += [H(2, "6. Expected Value & Kriteria Keputusan (KUK 2.4)"),
          P("Kriteria keputusan materi: " + "; ".join(f"**{k}**: {v}" for k, v in R["decision_criteria"].items())),
          T(["Berkas expected", "Target", "Jenis", "Status", "Cocok/total", "SHA-256"],
            [[m["id"], m["target"], m["kind"], m["status"], f"{m['summary']['matched']}/{m['summary']['total']}" if m.get("summary") else "—", m["sha256"][:16] + "…"]
             for m in sorted(ex["manifest"], key=lambda m: (m["status"] != "COCOK_VEKTOR_RESMI", m["id"])) if m["kind"] in ("deterministic", "negative") or m["status"] != "KRITERIA"][:120])]
    # 7
    D += [H(2, "7. Pemetaan Parameter → Kemungkinan Hasil Uji (KUK 2.5)"),
          P("Kategori hasil: " + "; ".join(f"**{k}** — {v}" for k, v in R["outcome_categories"].items())),
          T(["Kondisi temuan", "Parameter pemicu", "Kemungkinan hasil", "Tindak lanjut", "Sumber"],
            [[o["condition"], o["trigger"], o["outcome"], o["action"], o["source"]] for o in R["outcome_map"]])]
    # 8
    D += [H(2, "8. Matriks Keterlacakan"),
          T(["Kebutuhan", "Uraian", "Metode UK-1", "Skenario", "Test case", "Area", "Standar"],
            [[r["requirement_id"], r["requirement"], r["method"], r["scenario"], r["test_case"], r["area"], r["standards"]] for r in R["traceability"][:400]])]
    if pf["testing_model"] == 4:
        D += [H(2, "Tata Kelola Pengujian (model 4 tingkat)"),
              T(["Tingkat", "Pelaksana", "Kriteria masuk", "Kriteria keluar"], [[k, v.get("executor"), v.get("entry"), v.get("exit")] for k, v in R["governance"].items()])]
    # 9
    D += [H(2, "9. Lampiran: Peta KUK → Bab → File Kode"),
          T(["KUK", "Deskripsi", "Bab", "File", "Test"], [[k["kuk"], k["desc"], k["section"], ", ".join(k["files"]), k["tests"]] for k in KUK_MAP])]
    return D


KUK_MAP = [
    {"kuk": "1.1", "desc": "Metode ditelaah sesuai kebutuhan pengujian", "section": "Bab 2", "files": ["uk2_skenario/method_review.py"], "tests": "tests/test_uk2.py::test_kuk_1_1_*"},
    {"kuk": "1.2", "desc": "Parameter ditelaah sesuai metode", "section": "Bab 1, 2", "files": ["uk2_skenario/param_space.py"], "tests": "tests/test_uk2.py::test_kuk_1_2_*"},
    {"kuk": "2.1", "desc": "Skenario didesain dari hasil telaah", "section": "Bab 3", "files": ["uk2_skenario/designer.py", "uk2_skenario/data/recipes.yaml", "uk2_skenario/data/templates.yaml"], "tests": "tests/test_uk2.py::test_kuk_2_1_*"},
    {"kuk": "2.2", "desc": "Skenario diverifikasi terhadap spesifikasi desain (aspek kritis)", "section": "Bab 4", "files": ["uk2_skenario/verifier.py"], "tests": "tests/test_uk2.py::test_kuk_2_2_*"},
    {"kuk": "2.3", "desc": "Identifikasi test case", "section": "Bab 5", "files": ["uk2_skenario/testcase.py", "uk2_skenario/schema/uk2_skenario.schema.json"], "tests": "tests/test_uk2.py::test_kuk_2_3_*"},
    {"kuk": "2.4", "desc": "Expected value", "section": "Bab 6", "files": ["uk2_skenario/expected.py", "uk2_skenario/data/vectors/"], "tests": "tests/test_uk2.py::test_kuk_2_4_*"},
    {"kuk": "2.5", "desc": "Kompilasi skenario", "section": "Bab 7, 8", "files": ["uk2_skenario/compiler.py"], "tests": "tests/test_uk2.py::test_kuk_2_5_*"},
]


def validate_schema(data: dict) -> list:
    try:
        import jsonschema
    except ImportError:
        return ["jsonschema tidak terpasang — validasi dilewati"]
    schema = json.loads(SCHEMA.read_text())
    errs = sorted(jsonschema.Draft202012Validator(schema).iter_errors(data), key=lambda e: list(e.path))
    return [f"{'/'.join(map(str, e.path))}: {e.message}" for e in errs[:20]]


def write_all(R: dict, out: Path, docx: bool = True) -> dict:
    out.mkdir(parents=True, exist_ok=True)
    files = {}
    R["schema_errors"] = validate_schema(R)
    (out / OUT_JSON).write_text(json.dumps(R, indent=1, ensure_ascii=False, default=str), encoding="utf-8")
    files["json"] = out / OUT_JSON
    doc = build_document(R)
    (out / OUT_MD).write_text(to_markdown(doc), encoding="utf-8")
    files["md"] = out / OUT_MD
    if docx and to_docx(doc, out / OUT_DOCX):
        files["docx"] = out / OUT_DOCX
    cols = ["id", "tier", "layer", "area", "category", "title", "targets", "parameters", "prerequisites", "steps", "pass",
            "methods", "standards", "priority", "executor", "expected_kind", "expected_status", "runner_action"]
    rows = [[t["id"], t["tier"], t["layer"], t.get("area") or "", t["category"], t["title"], "; ".join(t["targets"]),
             "; ".join(f"{p['name']}={p['value']}({p['expect']})" for p in t["parameters"]), " | ".join(t["prerequisites"]),
             " ".join(t["steps"]), t["pass"], "; ".join(t["methods"]), "; ".join(t["standards"]), t["priority"], t["executor"],
             t["expected_kind"], "; ".join(t["expected_status"]), t["runner"]["action"]] for t in R["test_cases"]]
    with open(out / "test_cases.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(cols)
        w.writerows(rows)
    files["csv"] = out / "test_cases.csv"
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = "Test case"
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F5EFF")
        for r in rows:
            ws.append(r)
        ws2 = wb.create_sheet("Keterlacakan")
        ws2.append(list(R["traceability"][0].keys()) if R["traceability"] else [])
        for r in R["traceability"]:
            ws2.append(list(r.values()))
        ws3 = wb.create_sheet("Pemetaan hasil")
        ws3.append(["condition", "trigger", "outcome", "action", "source"])
        for o in R["outcome_map"]:
            ws3.append([o["condition"], o["trigger"], o["outcome"], o["action"], o["source"]])
        wb.save(out / "test_cases.xlsx")
        files["xlsx"] = out / "test_cases.xlsx"
    except ImportError:
        pass
    with open(out / "matriks_keterlacakan.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(R["traceability"][0].keys()) if R["traceability"] else ["requirement_id"])
        w.writeheader()
        w.writerows(R["traceability"])
    files["trace"] = out / "matriks_keterlacakan.csv"
    (out / "laporan_verifikasi.md").write_text(V.to_markdown(R["verification"], R["profile"]["product"]["name"]), encoding="utf-8")
    files["verif"] = out / "laporan_verifikasi.md"
    return files
