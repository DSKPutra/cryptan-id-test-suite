"""Ekstrak tabel skenario dari materi UK-2 (PPTX) → templat YAML di uk2_skenario/data/.

  python scripts/extract_uk2_templates.py [--materi docs/materi]

Sumber (hak cipta Politeknik Siber dan Sandi Negara, hanya disimpan lokal):
  - UK-2_Skenario_Pengujian_Produk_Kriptografi.pptx  (6 kategori, contoh ECDSA & TLS 1.3)
  - UK-2_Skenario_Uji_Modul_Kriptografi_ISO19790.pptx (HSM-X, 4 tingkat, 11 area)
Keluaran YAML adalah turunan terstruktur (tabel), dengan atribusi sumber per bagian.
"""
import argparse
import re
from pathlib import Path

import yaml
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "uk2_skenario" / "data"
F1 = "UK-2_Skenario_Pengujian_Produk_Kriptografi.pptx"
F2 = "UK-2_Skenario_Uji_Modul_Kriptografi_ISO19790.pptx"
ATTR1 = "Materi UK-2 'Penetapan Skenario Pengujian Produk Kriptografi' (Dr. Santi Indarjani, Poltek SSN)"
ATTR2 = "Materi UK-2 'Penetapan Skenario Pengujian Modul Kriptografi Berbasis ISO/IEC 19790' (Dr. Santi Indarjani, Poltek SSN)"

LAYER = {"Uji Kesesuaian": "K", "Uji Keamanan": "S", "Uji Implementasi": "I"}
CATS = [  # (key, kode ID, label, slide parameter, slide matriks)
    ("block", "BC", "Block Cipher", 6, 7),
    ("stream", "SC", "Stream Cipher", 8, 9),
    ("hash", "HF", "Fungsi Hash", 10, 11),
    ("pke", "PK", "Public Key Encryption", 12, 13),
    ("signature", "DS", "Digital Signature", 14, 15),
    ("protocol", "PR", "Protokol Kriptografi", 16, 17),
]


def clean(s: str) -> str:
    return re.sub(r"\s+", " ", s.replace(" ", " ")).strip()


def tables(slide):
    return [[[clean(c.text) for c in r.cells] for r in sh.table.rows] for sh in slide.shapes if sh.has_table]


def texts(slide):
    out = []
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            out.append(sh.text_frame.text.strip())
    return out


def notes_after_table(slide, skip_titles=2):
    t = [l for x in texts(slide)[skip_titles:] for l in x.splitlines() if l.strip()]
    return [clean(x) for x in t if x not in ("Catatan kritis",) and not x.startswith("Setiap sel")]


def kv_block(text: str) -> dict:
    lines = [clean(l) for l in text.splitlines() if l.strip()]
    return {lines[i]: lines[i + 1] for i in range(0, len(lines) - 1, 2)}


def tc_table(rows, level=None, layer=None, area_col=False):
    head, out = rows[0], []
    for r in rows[1:]:
        d = dict(zip(head, r))
        tc = {"id": d["ID"], "params": d["Parameter & Nilai Uji"], "procedure": d["Prosedur"], "pass": d["Kriteria Lulus"]}
        if "Tingkat" in d:
            tc["level"] = d["Tingkat"]
        if area_col:
            tc["area"] = d["Area"]
        if level:
            tc["level"] = level
        tc["layer"] = layer or d["ID"].split("-")[1]
        out.append(tc)
    return out


def outcome_table(rows):
    return [{"condition": r[0], "trigger": r[1], "outcome": r[2], "action": r[3]} for r in rows[1:]]


def extract_f1(path: Path) -> dict:
    p = Presentation(path)
    S = lambda n: p.slides[n - 1]
    # Slide 3 — matriks umum 3 lapis × 3 tingkat
    t = texts(S(3))[5:]
    framework = {}
    for i in range(0, len(t), 4):
        name, desc = t[i].split("\n", 1)
        framework[name] = {"desc": clean(desc), "K": clean(t[i + 1]), "S": clean(t[i + 2]), "I": clean(t[i + 3])}
    # Slide 4 — klasifikasi parameter & teknik
    t4 = texts(S(4))
    classes, techniques = [], []
    ti = t4.index("Teknik penetapan nilai uji")
    for i, x in enumerate(t4[:ti]):
        if x in ("1", "2", "3", "4", "5"):
            a, b = t4[i + 1].split("\n", 1)
            classes.append({"no": int(x), "class": clean(a), "examples": clean(b)})
    for x in t4[ti + 1:]:
        a, b = x.split("\n", 1)
        techniques.append({"technique": clean(a), "desc": clean(b)})
    # Slide 5 — kategori hasil & kriteria keputusan
    t5 = texts(S(5))
    outcomes = {}
    for name in ("Memenuhi", "Memenuhi dengan Catatan", "Tidak Memenuhi", "Inkonklusif"):
        i = t5.index(name)
        outcomes[name] = clean(t5[i + 1])
    decision = {}
    for x in t5[t5.index("Kriteria keputusan") + 1:]:
        a, b = x.split("\n", 1)
        decision[clean(a)] = clean(b)
    cats = {}
    for key, code, label, sp, sm in CATS:
        prm = tables(S(sp))[0]
        mat = tables(S(sm))[0]
        cats[key] = {
            "code": code, "label": label, "source_slides": [sp, sm],
            "params": [{"parameter": r[0], "variations": r[1], "impact": r[2]} for r in prm[1:]],
            "notes": notes_after_table(S(sp)),
            "matrix": {r[0]: {"K": r[1], "S": r[2], "I": r[3]} for r in mat[1:]},
        }
    # Slide 18 — lintas-algoritma
    t18 = texts(S(18))
    chain = [clean(x) for x in t18[2:8]]
    cross = [{"theme": clean(x.split("\n")[0]), "desc": clean(x.split("\n", 1)[1])} for x in t18 if "\n" in x]

    def example(sp, sk, ss, si, sm, code):
        prof = kv_block(texts(S(sp))[3])
        ps = tables(S(sp))[0]
        tcs = []
        for sl, lay in ((sk, "K"), (ss, "S"), (si, "I")):
            tcs += tc_table(tables(S(sl))[0], layer=lay)
        return {"source": ATTR1, "source_slides": [sp, sk, ss, si, sm], "code": code, "profile": prof,
                "param_space": [{"parameter": r[0], "status": r[1], "values": r[2]} for r in ps[1:]],
                "test_cases": tcs, "outcome_map": outcome_table(tables(S(sm))[0])}

    return {
        "templates": {"source": ATTR1, "framework_3x3": framework, "param_classes": classes, "techniques": techniques,
                      "outcome_categories": outcomes, "decision_criteria": decision, "categories": cats,
                      "cross_algorithm": {"chain": chain, "themes": cross}},
        "ecdsa": example(20, 21, 22, 23, 24, "DS"),
        "tls13": example(25, 26, 27, 28, 29, "PR"),
    }


def extract_f2(path: Path) -> dict:
    p = Presentation(path)
    S = lambda n: p.slides[n - 1]
    t3 = texts(S(3))
    prof = kv_block(t3[2])
    i0 = t3.index("Batas Kriptografis")
    boundary = [clean(x) for x in t3[i0 + 1:] if x not in ("Data input", "Data output", "Control input", "Status output", "Power")]
    ports = [x for x in t3 if x in ("Data input", "Data output", "Control input", "Status output", "Power")]
    # Slide 4 — 11 area + lapis dominan
    t4 = texts(S(4))[2:]
    t4 = t4[:t4.index("Uji Kesesuaian") - 1]          # buang legenda (K/S/I + keterangan)
    areas, cur = [], None
    for x in t4:
        if re.fullmatch(r"7\.\d+", x):
            cur = {"area": x, "name": None, "dominant": []}
            areas.append(cur)
        elif cur and cur["name"] is None:
            cur["name"] = clean(x)
        elif cur and x in ("K", "S", "I") and len(cur["dominant"]) < 3:
            cur["dominant"].append(x)
        elif x.startswith("Uji") or x.startswith("Label"):
            break
    # Slide 5 — matriks 3 lapis × 4 tingkat
    t5 = texts(S(5))[5:]
    fw = {}
    for i in range(0, len(t5), 4):
        name, desc = t5[i].split("\n", 1)
        fw[clean(name)] = {"desc": clean(desc), "K": clean(t5[i + 1]), "S": clean(t5[i + 2]), "I": clean(t5[i + 3])}
    ps = tables(S(6))[0]
    # Slide 7 — tata kelola
    t7 = texts(S(7))[2:]
    gov = {}
    keymap = {"Pelaksana": "executor", "Kriteria masuk": "entry", "Kriteria keluar": "exit"}
    for i in range(0, len(t7) - 1, 2):
        name = clean(re.sub(r"^\d\s+", "", t7[i]))
        gov[name] = {keymap[k]: v for k, v in kv_block(t7[i + 1]).items()}
    tcs = []
    levels = [("Unit", "U", (8, 9, 10)), ("Integrasi", "I", (11, 12, 13)), ("Sistem", "S", (14, 15, 16)), ("UAT", "A", (17, 18, 19))]
    for lname, _, slides in levels:
        for sl, lay in zip(slides, ("K", "S", "I")):
            tcs += tc_table(tables(S(sl))[0], level=lname, layer=lay, area_col=True)
    cov_rows = tables(S(20))[0]
    coverage = {}
    for r in cov_rows[1:]:
        a = r[0].split()[0]
        coverage[a] = {lv: ([] if c in ("—", "-", "") else [x.strip() for x in c.split(",")])
                       for lv, c in zip(("Unit", "Integrasi", "Sistem", "UAT"), r[1:])}
    return {"hsm": {"source": ATTR2, "code": "HSM", "model": "4 tingkat", "security_level": 3,
                    "profile": prof, "boundary_components": boundary, "ports": ports,
                    "param_space": [{"parameter": r[0], "values": r[1], "area": r[2], "impact": r[3]} for r in ps[1:]],
                    "test_cases": tcs, "coverage": coverage, "outcome_map": outcome_table(tables(S(21))[0])},
            "iso": {"source": ATTR2, "areas": areas, "framework_3x4": fw, "governance": gov,
                    "standards": ["ISO/IEC 19790:2012", "ISO/IEC 24759:2017", "ISO/IEC 18367", "ISO/IEC 17825"]}}


def dump(obj, name):
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / name).write_text("# DIHASILKAN oleh scripts/extract_uk2_templates.py dari materi PPTX — jangan disunting manual\n"
                            + yaml.safe_dump(obj, sort_keys=False, allow_unicode=True, width=140), encoding="utf-8")
    print("→", OUT / name)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--materi", default=str(ROOT / "docs" / "materi"))
    a = ap.parse_args()
    m = Path(a.materi)
    r1, r2 = extract_f1(m / F1), extract_f2(m / F2)
    dump(r1["templates"], "templates.yaml")
    (OUT / "examples").mkdir(parents=True, exist_ok=True)
    dump(r1["ecdsa"], "examples/ecdsa_p256_token.yaml")
    dump(r1["tls13"], "examples/tls13_gateway.yaml")
    dump(r2["hsm"], "examples/hsm_x_sl3.yaml")
    dump(r2["iso"], "iso19790.yaml")
