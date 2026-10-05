"""Keluaran HTML (pratinjau) dan XLSX (daftar + hasil) dari data yang sama dengan PDF (results.json)."""
from html import escape as h
from pathlib import Path

from .builder import _fmt, _params
from .model import Model, load, tanggal_id
from .template import THEME

COL = THEME["colors"]


def _badge(status, overall=False):
    pal = COL["overall"] if overall else COL["status"]
    bg = pal.get(status, COL["muted"])
    fg = COL["text"] if status == "TIDAK_BERLAKU" else "#fff"
    return f'<span class="b" style="background:{bg};color:{fg}">{h(status.replace("_", " "))}</span>'


def build_html(res=None, out=None, sources=None, primitives=None, only_tested=False, penyusun=None, ringkas=False) -> Path:
    res = res or load()
    m = Model(res, sources, primitives, only_tested, penyusun)
    k, ident = m.kpi(), THEME["identity"]
    acc = COL["accent"]
    parts = [f"""<!doctype html><html lang="id"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1"><title>{h(THEME['short_title'])}</title>
<style>
:root{{--acc:{acc};--zebra:{COL['zebra']};--text:{COL['text']};--muted:{COL['muted']};--grid:{COL['grid']}}}
body{{font-family:"DejaVu Sans",system-ui,sans-serif;color:var(--text);margin:0;background:#fff;font-size:14px}}
header{{background:var(--acc);color:#fff;padding:14px 24px;display:flex;justify-content:space-between;gap:12px;flex-wrap:wrap}}
header .cls{{background:#C62828;padding:2px 10px;border-radius:3px;font-weight:700}}
main{{max-width:1180px;margin:0 auto;padding:16px}}
h1{{color:var(--acc);font-size:24px}} h2{{color:var(--acc);border-bottom:2px solid var(--acc);padding-bottom:4px;margin-top:32px}}
.kpi{{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:10px}}
.kpi div{{background:#E3EFED;border:1px solid var(--acc);border-radius:6px;padding:12px;text-align:center}}
.kpi b{{display:block;font-size:28px;color:var(--acc)}} .kpi span{{color:var(--muted);font-size:12px}}
.wrap{{overflow-x:auto}} table{{border-collapse:collapse;width:100%;font-size:12px;margin:8px 0 16px}}
th{{background:var(--acc);color:#fff;text-align:left;padding:5px;position:sticky;top:0}}
td{{border:1px solid var(--grid);padding:4px 5px;vertical-align:top}} tr:nth-child(even) td{{background:var(--zebra)}}
.b{{display:inline-block;padding:1px 6px;border-radius:3px;font-weight:700;font-size:11px;white-space:nowrap}}
code{{font-family:"DejaVu Sans Mono",monospace;font-size:11.5px}} .muted{{color:var(--muted);font-size:12px}}
footer{{color:var(--muted);font-size:12px;padding:16px 24px;border-top:1px solid var(--grid)}}
</style></head><body>
<header><span>{h(THEME['title'])}</span><span class="cls">{h(ident['klasifikasi'])}</span></header><main>
<h1>{h(THEME['title'])}</h1><p class="muted">{h(THEME['subtitle'])} · Penyusun: {h(penyusun or ident['penyusun_default'])} ·
{h(ident['instansi'])} · {h(tanggal_id(res['generated_at']))} · mode {h(res['mode'])} · seed {res['seed']}</p>
<h2>1 Ringkasan eksekutif</h2><div class="kpi">
<div><b>{k['terdaftar']}</b><span>algoritme terdaftar</span></div><div><b>{k['diuji']}</b><span>diuji</span></div>
<div><b>{k['lulus_semua']}</b><span>lulus semua uji</span></div><div><b>{k['punya_temuan']}</b><span>punya temuan</span></div>
<div><b>{k['tidak_dapat_diuji']}</b><span>tidak dapat diuji</span></div></div><ul>"""]
    parts += [f"<li>{h(x)}</li>" for x in m.top_findings()]
    parts.append("</ul><h2>3 Daftar algoritme standar</h2>")
    for b in m.bodies:
        parts.append(f"<h3>{h(m.body_labels[b])}</h3><div class='wrap'><table><tr><th>No</th><th>ID</th><th>Primitif</th>"
                     "<th>Varian</th><th>Kunci</th><th>Kekuatan</th><th>Status NIST</th><th>Dokumen</th><th>Status uji</th></tr>")
        i = 0
        for prim, rows in m.by_body(b).items():
            for a in rows:
                i += 1
                parts.append(f"<tr><td>{i}</td><td><code>{h(a['id'])}</code></td><td>{h(m.prim_labels[prim])}</td>"
                             f"<td>{h(a['variant'])}</td><td>{h(_fmt(a['key_bits']))}</td><td>{h(_fmt(a['security_strength_bits']))}</td>"
                             f"<td>{h(_fmt(a['status_nist']))}</td><td>{h(', '.join(m.docs_for_body(a, b)))}</td>"
                             f"<td>{_badge(a['overall'], True)}</td></tr>")
        parts.append("</table></div>")
    parts.append("<h2>4 Matriks hasil uji</h2>")
    for prim, lbl in m.primitives_present().items():
        parts.append(f"<h3>{h(lbl)}</h3><div class='wrap'><table><tr><th>Algoritme</th>" +
                     "".join(f"<th>{t}</th>" for t in m.test_ids) + "<th>Keseluruhan</th></tr>")
        for a in (x for x in m.algs if x["primitive"] == prim):
            st = {t["plugin"]: t["status"] for t in a["tests"]}
            parts.append(f"<tr><td><code>{h(a['id'])}</code></td>" + "".join(f"<td>{_badge(st.get(t, 'TIDAK_BERLAKU'))}</td>"
                                                                           for t in m.test_ids) +
                         f"<td>{_badge(a['overall'], True)}</td></tr>")
        parts.append("</table></div>")
    if not ringkas:
        parts.append("<h2>5 Hasil rinci</h2><div class='wrap'><table><tr><th>Uji</th><th>Status</th><th>Parameter</th>"
                     "<th>Hasil aktual</th><th>Kriteria</th><th>Evidence</th></tr>")
        for a in m.algs:
            for t in a["tests"]:
                if t["status"] == "TIDAK_BERLAKU":
                    continue
                parts.append(f"<tr><td><code>{h(t['test_id'])}</code></td><td>{_badge(t['status'])}</td>"
                             f"<td>{h(_params(t.get('parameters')))}</td><td>{h(t.get('actual') or t.get('reason') or '')}</td>"
                             f"<td>{h(t.get('criteria') or '')}</td><td><code>{h(t.get('evidence_id') or '—')}</code></td></tr>")
        parts.append("</table></div>")
    parts.append("<h2>6 Temuan</h2><div class='wrap'><table><tr><th>Algoritme</th><th>Uji</th><th>Status</th><th>Hasil &amp; bukti</th></tr>")
    for a, t, det in m.failing():
        parts.append(f"<tr><td><code>{h(a['id'])}</code></td><td>{t['plugin']}</td><td>{_badge(t['status'])}</td>"
                     f"<td>{h((t.get('actual') or '') + (' — ' + det if det else ''))}</td></tr>")
    parts.append("</table></div><h2>Lampiran C — PERLU_VERIFIKASI</h2><div class='wrap'><table><tr><th>Butir</th><th>Jenis</th><th>Rincian</th></tr>")
    parts += [f"<tr><td><code>{h(x['item'])}</code></td><td>{h(x['jenis'])}</td><td>{h(x['detail'])}</td></tr>" for x in res["perlu_verifikasi"]]
    parts.append(f"</table></div></main><footer>{h(ident['copyright'])} · dihasilkan dari results.json · "
                 f"{h(ident['klasifikasi'])}</footer></body></html>")
    path = Path(out) if out else m.filename("html")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(parts), encoding="utf-8")
    return path


def build_xlsx(res=None, out=None, sources=None, primitives=None, only_tested=False, penyusun=None) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    res = res or load()
    m = Model(res, sources, primitives, only_tested, penyusun)
    wb = Workbook()
    hdr_fill = PatternFill("solid", fgColor=COL["accent"].lstrip("#"))
    hdr_font = Font(bold=True, color="FFFFFF")

    def sheet(ws, header, rows, widths):
        ws.append(header)
        for c in ws[1]:
            c.fill, c.font = hdr_fill, hdr_font
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=True)
                pal = COL["status"].get(c.value) or COL["overall"].get(c.value)
                if isinstance(c.value, str) and pal:
                    c.fill = PatternFill("solid", fgColor=pal.lstrip("#"))
                    c.font = Font(bold=True, color="1F2933" if c.value == "TIDAK_BERLAKU" else "FFFFFF")

    k = m.kpi()
    ws = wb.active
    ws.title = "Ringkasan"
    ident = THEME["identity"]
    sheet(ws, ["Butir", "Nilai"], [["Judul", THEME["title"]], ["Klasifikasi", ident["klasifikasi"]],
                                   ["Penyusun", penyusun or ident["penyusun_default"]], ["Instansi", ident["instansi"]],
                                   ["Tanggal", tanggal_id(res["generated_at"])], ["Mode", res["mode"]], ["Seed", res["seed"]]] +
          [[f"KPI: {a.replace('_', ' ')}", b] for a, b in k.items()] +
          [[f"Jumlah hasil uji: {a}", b] for a, b in m.status_counts().items()] + [["Hak cipta", ident["copyright"]]], [34, 70])
    sheet(wb.create_sheet("Daftar algoritme"),
          ["ID", "Primitif", "Keluarga", "Varian", "Kunci (bit)", "Blok/keluaran (bit)", "Kekuatan (bit)", "Status NIST",
           "source_body", "source_doc", "Rentan kuantum", "Backend", "Status keseluruhan"],
          [[a["id"], m.prim_labels.get(a["primitive"], a["primitive"]), a["family"], a["variant"], _fmt(a["key_bits"]),
            _fmt(a["block_or_output_bits"]), _fmt(a["security_strength_bits"]), _fmt(a["status_nist"]), ", ".join(a["source_body"]),
            ", ".join(a["source_doc"]), "ya" if a.get("quantum_vulnerable") else "tidak", ", ".join(a["backends"]), a["overall"]]
           for a in m.algs], [24, 20, 14, 16, 10, 12, 12, 14, 20, 36, 10, 22, 20])
    sheet(wb.create_sheet("Matriks"), ["Algoritme"] + m.test_ids + ["Keseluruhan"],
          [[a["id"]] + [{t["plugin"]: t["status"] for t in a["tests"]}.get(x, "TIDAK_BERLAKU") for x in m.test_ids] + [a["overall"]]
           for a in m.algs], [24] + [17] * len(m.test_ids) + [20])
    sheet(wb.create_sheet("Hasil uji"), ["ID uji", "Algoritme", "Uji", "Nama", "Status", "Label", "Parameter", "Hasil aktual",
                                         "Kriteria", "Alasan", "Durasi (s)", "Evidence ID", "Sumber"],
          [[t["test_id"], a["id"], t["plugin"], t["name"], t["status"], t.get("label", ""), _params(t.get("parameters")),
            t.get("actual", ""), t.get("criteria", ""), t.get("reason", ""), t.get("duration_s"), t.get("evidence_id") or "",
            t.get("source", "")] for a in m.algs for t in a["tests"]],
          [26, 20, 7, 30, 18, 16, 40, 40, 26, 30, 9, 15, 26])
    sheet(wb.create_sheet("Vektor"), ["Berkas", "Sumber", "URL", "Lisensi", "Uji subset", "Uji penuh", "SHA-256"],
          [[v["file"], v["source"], v.get("url", ""), v.get("license", ""), v.get("tests_subset"), v.get("tests_full"), v["sha256"]]
           for v in res["vectors"]], [40, 34, 50, 12, 10, 10, 66])
    sheet(wb.create_sheet("PERLU_VERIFIKASI"), ["Butir", "Jenis", "Rincian"],
          [[x["item"], x["jenis"], x["detail"]] for x in res["perlu_verifikasi"]], [30, 26, 70])
    path = Path(out) if out else m.filename("xlsx")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
