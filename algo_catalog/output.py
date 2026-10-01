"""Keluaran outputs/algo_catalog/: algorithms_under_test.yaml, Daftar_Algoritma_Uji.md,
.csv, .xlsx + peringatan otomatis, dan tautan otomatis ke config/product_profile.yaml.
"""
import csv
import datetime as dt
import re
from pathlib import Path

import yaml

from .schema import PRIMITIVES, UK1_PRIMITIVE
from .selection import summarize, warnings_for

ROOT = Path(__file__).resolve().parents[1]
PROFILE = ROOT / "config" / "product_profile.yaml"
YAML_NAME = "algorithms_under_test.yaml"
MD_NAME = "Daftar_Algoritma_Uji.md"
CSV_NAME = "Daftar_Algoritma_Uji.csv"
XLSX_NAME = "Daftar_Algoritma_Uji.xlsx"
COLUMNS = ["No", "Primitif", "Algoritma", "Varian", "Panjang kunci", "Security strength", "Status NIST",
           "Standar", "Sumber input", "Bukti"]
INPUT_LABEL = {"link": "link", "file": "file", "dropdown": "dropdown", "edit": "suntingan"}


def _uk1_profiles() -> dict:
    """catalog_id → berkas profil UK-1 (config/profiles/*.yaml)."""
    out = {}
    for f in sorted((ROOT / "config" / "profiles").glob("*.yaml")):
        d = yaml.safe_load(f.read_text(encoding="utf-8")) or {}
        cid = (d.get("algorithm") or {}).get("catalog_id")
        if cid:
            out[cid.upper()] = f"config/profiles/{f.name}"
    return out


def _strength(r):
    s = r["security_strength_bits"]
    cat = r.get("security_category")
    txt = "—" if s is None else (f"{s} bit" if isinstance(s, int) else str(s))
    return txt + (f" (kategori {cat})" if cat not in (None, "PERLU_VERIFIKASI") else "")


def _sources(r):
    return ", ".join(sorted({INPUT_LABEL.get(i["type"], i["type"]) for i in r["inputs"]}))


def _evidence(r, n=1):
    ev = []
    for i in r["inputs"][:n]:
        loc = f"{i.get('ref', '')}" + (f":{i['line']}" if i.get("line") else "")
        ev.append(f"[{loc}] {i.get('evidence', '')}"[:160])
    return " · ".join(ev)


def table_rows(rows):
    return [[n, PRIMITIVES[r["primitive"]], r["family"], r["variant"], "—" if r["key_bits"] is None else r["key_bits"],
             _strength(r), r["status_nist"], ", ".join(r["standards"]), _sources(r), _evidence(r)]
            for n, r in enumerate(rows, 1)]


def build_document(rows, product: str, pending=()) -> dict:
    uk1 = _uk1_profiles()
    warns = [w for r in rows for w in warnings_for(r)]
    items = []
    for r in rows:
        it = {k: r[k] for k in ("id", "entry_id", "primitive", "family", "variant", "key_bits", "block_or_output_bits",
                                "security_strength_bits", "security_category", "status_nist", "standards", "quantum_vulnerable")}
        it["uk1_primitive"] = UK1_PRIMITIVE[r["primitive"]]
        it["uk1_profile"] = uk1.get(r["id"].upper())
        it["confidence"] = r["confidence"]
        it["inputs"] = [{k: v for k, v in i.items() if v is not None} for i in r["inputs"]]
        it["warnings"] = [w["code"] for w in warnings_for(r)]
        items.append(it)
    return {
        "schema": "cryptan.algo_catalog.algorithms_under_test.v1",
        "generated_at": dt.datetime.now().astimezone().isoformat(timespec="seconds"),
        "product": product,
        "kuk": "UK-1 KUK 1.1 — informasi desain produk (daftar algoritma yang diuji)",
        "summary": summarize(rows),
        "uk1_coverage": {
            "with_profile": sorted({i["uk1_profile"] for i in items if i["uk1_profile"]}),
            "without_profile": [i["id"] for i in items if not i["uk1_profile"]],
            "note": "Kombinasi tanpa profil UK-1 dapat dibuat dengan menyalin profil primitif yang sama di config/profiles/",
        },
        "warnings": warns,
        "pending_unconfirmed": [{"id": p["id"], "confidence": p["confidence"], "evidence": _evidence(p)} for p in pending],
        "items": items,
    }


def to_markdown(doc: dict, rows) -> str:
    s = doc["summary"]
    L = [f"# Daftar Algoritma yang Diuji — {doc['product']}", "",
         f"*Cryptan.ID Test Suite · modul `algo_catalog` · {doc['kuk']} · dibangkitkan {doc['generated_at']}*", "",
         "## Ringkasan", "",
         f"Total **{s['total_combinations']}** kombinasi algoritma × varian × panjang kunci.", "",
         "| Primitif | Jumlah kombinasi |", "|---|---|"]
    L += [f"| {PRIMITIVES[p]} | {n} |" for p, n in s["by_primitive"].items()]
    L += [f"| **Total** | **{s['total_combinations']}** |", "",
          "Sumber input: " + ", ".join(f"{k} = {v}" for k, v in s["by_input"].items() if v), "",
          "## Daftar algoritma", "", "| " + " | ".join(COLUMNS) + " |", "|" + "---|" * len(COLUMNS)]
    for row in table_rows(rows):
        L.append("| " + " | ".join(str(c).replace("|", "\\|") for c in row) + " |")
    L += ["", "## Peringatan otomatis", ""]
    if doc["warnings"]:
        L += ["| Tingkat | Kode | Pesan |", "|---|---|---|"]
        L += [f"| {w['level']} | {w['code']} | {w['message']} |" for w in doc["warnings"]]
    else:
        L.append("Tidak ada peringatan.")
    cov = doc["uk1_coverage"]
    L += ["", "## Tautan ke UK-1", "",
          f"Berkas ini ditautkan ke `config/product_profile.yaml` (`algorithms_under_test`).",
          f"Profil UK-1 tersedia untuk: {', '.join(cov['with_profile']) or '—'}.",
          f"Kombinasi tanpa profil UK-1 ({len(cov['without_profile'])}): {', '.join(cov['without_profile'][:40]) or '—'}"
          + (" …" if len(cov["without_profile"]) > 40 else ""), ""]
    if doc["pending_unconfirmed"]:
        L += ["## Kandidat belum dikonfirmasi (tidak masuk daftar akhir)", "", "| ID | Keyakinan | Bukti |", "|---|---|---|"]
        L += [f"| {p['id']} | {p['confidence']} | {p['evidence'].replace('|', '/')} |" for p in doc["pending_unconfirmed"]]
    return "\n".join(L) + "\n"


def link_profile(yaml_path: Path, profile_path: Path = PROFILE) -> bool:
    """Tautkan otomatis: set `algorithms_under_test: <path>` di product_profile.yaml (komentar dipertahankan)."""
    if not profile_path.exists():
        return False
    rel = Path("..") / yaml_path.resolve().relative_to(ROOT) if yaml_path.resolve().is_relative_to(ROOT) else yaml_path.resolve()
    line = f"algorithms_under_test: {rel.as_posix()}   # dihasilkan algo_catalog (KUK 1.1)\n"
    text = profile_path.read_text(encoding="utf-8")
    if re.search(r"^algorithms_under_test:.*$", text, re.M):
        text = re.sub(r"^algorithms_under_test:.*\n?", line, text, flags=re.M)
    else:
        text = text.replace("\n# Profil algoritma objek praktikum", "\n" + line + "\n# Profil algoritma objek praktikum", 1) \
            if "# Profil algoritma objek praktikum" in text else text.rstrip("\n") + "\n" + line
    profile_path.write_text(text, encoding="utf-8")
    return True


def write_all(rows, out_dir: Path, product: str = "Cryptan.ID SecureLib", pending=(), link: bool = True) -> dict:
    out_dir.mkdir(parents=True, exist_ok=True)
    doc = build_document(rows, product, pending)
    files = {}
    p = out_dir / YAML_NAME
    p.write_text("# Dihasilkan oleh: python -m algo_catalog select …  (jangan disunting manual)\n"
                 + yaml.safe_dump(doc, sort_keys=False, allow_unicode=True, width=120), encoding="utf-8")
    files["yaml"] = p
    files["json"] = out_dir / "algorithms_under_test.json"          # salinan untuk dashboard web
    files["json"].write_text(__import__("json").dumps(doc, indent=1, ensure_ascii=False, default=str), encoding="utf-8")
    files["md"] = out_dir / MD_NAME
    files["md"].write_text(to_markdown(doc, rows), encoding="utf-8")
    trows = table_rows(rows)
    with open(out_dir / CSV_NAME, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(COLUMNS + ["ID kanonik", "Keyakinan", "Peringatan"])
        for r, row in zip(rows, trows):
            w.writerow(row + [r["id"], r["confidence"], ";".join(x["code"] for x in warnings_for(r))])
    files["csv"] = out_dir / CSV_NAME
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = "Daftar Algoritma"
        ws.append(COLUMNS + ["ID kanonik", "Keyakinan", "Peringatan"])
        for c in ws[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F5EFF")
        for r, row in zip(rows, trows):
            ws.append(row + [r["id"], r["confidence"], ";".join(x["code"] for x in warnings_for(r))])
        for col, wdt in zip("ABCDEFGHIJKLM", (5, 22, 12, 18, 13, 18, 16, 36, 16, 60, 24, 10, 30)):
            ws.column_dimensions[col].width = wdt
        ws2 = wb.create_sheet("Ringkasan")
        ws2.append(["Primitif", "Jumlah kombinasi"])
        for pr, n in doc["summary"]["by_primitive"].items():
            ws2.append([PRIMITIVES[pr], n])
        ws2.append(["Total", doc["summary"]["total_combinations"]])
        ws3 = wb.create_sheet("Peringatan")
        ws3.append(["Tingkat", "Kode", "Pesan"])
        for w in doc["warnings"]:
            ws3.append([w["level"], w["code"], w["message"]])
        wb.save(out_dir / XLSX_NAME)
        files["xlsx"] = out_dir / XLSX_NAME
    except ImportError:
        pass
    if link:
        files["profile_linked"] = PROFILE if link_profile(files["yaml"]) else None
    return {"files": files, "doc": doc}
