"""Opsi 2 (File) — ekstraksi teks per tipe berkas.

Didukung: .pdf (pypdf), .docx (python-docx), .xlsx (openpyxl), .csv, .json,
.yaml/.yml, .txt, .md, .html/.htm, serta source code (.c .h .cpp .py .java .go
.rs .js .ts .cs). Source code dibaca apa adanya agar pemanggilan pustaka
(mis. `EVP_aes_256_gcm`, `Cipher.getInstance("AES/GCM/NoPadding")`) terdeteksi.
"""
import csv
import io
import json
from html.parser import HTMLParser
from pathlib import Path

TEXT_EXT = {".txt", ".md", ".csv", ".log", ".ini", ".cfg", ".conf"}
CODE_EXT = {".c", ".h", ".cc", ".cpp", ".hpp", ".py", ".java", ".kt", ".go", ".rs", ".js", ".ts", ".cs", ".swift", ".rb", ".php"}
SUPPORTED = TEXT_EXT | CODE_EXT | {".pdf", ".docx", ".xlsx", ".json", ".yaml", ".yml", ".html", ".htm"}


class ExtractError(ValueError):
    pass


class _HTMLText(HTMLParser):
    SKIP = {"script", "style", "noscript", "svg"}

    def __init__(self):
        super().__init__()
        self.parts, self._skip = [], 0

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self._skip += 1
        elif tag in ("br", "p", "li", "tr", "h1", "h2", "h3", "h4", "div", "td", "th"):
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in self.SKIP and self._skip:
            self._skip -= 1

    def handle_data(self, data):
        if not self._skip:
            self.parts.append(data)


def html_to_text(html: str) -> str:
    p = _HTMLText()
    p.feed(html)
    lines = [" ".join(l.split()) for l in "".join(p.parts).splitlines()]
    return "\n".join(l for l in lines if l)


def pdf_to_text(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise ExtractError("pypdf belum terpasang (pip install pypdf)") from e
    r = PdfReader(io.BytesIO(data))
    return "\n".join((pg.extract_text() or "") for pg in r.pages)


def docx_to_text(data: bytes) -> str:
    import docx
    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def xlsx_to_text(data: bytes) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _flatten(obj, out):
    if isinstance(obj, dict):
        for k, v in obj.items():
            out.append(str(k))
            _flatten(v, out)
    elif isinstance(obj, (list, tuple)):
        for v in obj:
            _flatten(v, out)
    elif obj is not None:
        out.append(str(obj))


def struct_to_text(obj) -> str:
    out = []
    _flatten(obj, out)
    return "\n".join(out)


def bytes_to_text(name: str, data: bytes) -> str:
    """Ekstrak teks dari isi berkas berdasarkan ekstensi `name`."""
    ext = Path(name).suffix.lower()
    if ext not in SUPPORTED:
        raise ExtractError(f"tipe berkas tidak didukung: {ext or name}")
    if ext == ".pdf":
        return pdf_to_text(data)
    if ext == ".docx":
        return docx_to_text(data)
    if ext == ".xlsx":
        return xlsx_to_text(data)
    text = data.decode("utf-8", errors="replace")
    if ext in (".html", ".htm"):
        return html_to_text(text)
    if ext == ".json":
        try:
            return struct_to_text(json.loads(text)) + "\n" + text
        except json.JSONDecodeError:
            return text
    if ext in (".yaml", ".yml"):
        import yaml
        try:
            return struct_to_text(yaml.safe_load(text)) + "\n" + text
        except yaml.YAMLError:
            return text
    if ext == ".csv":
        rows = list(csv.reader(io.StringIO(text)))
        return "\n".join(" | ".join(r) for r in rows)
    return text


def file_to_text(path) -> str:
    p = Path(path)
    return bytes_to_text(p.name, p.read_bytes())
