"""Generator keluaran UK-1: JSON, Dokumen Penetapan Metode (Markdown + DOCX),
dan CSV matriks. Markdown & DOCX dibangun dari satu model dokumen yang sama.
"""
import csv
import json
from pathlib import Path

JSON_NAME = "uk1_penetapan_metode.json"
MD_NAME = "UK1_Dokumen_Penetapan_Metode.md"
DOCX_NAME = "UK1_Dokumen_Penetapan_Metode.docx"
CSV_MAP = "matriks_pemetaan.csv"
CSV_TRACE = "matriks_keterlacakan.csv"


# ---------------------------------------------------------------- model dokumen
def H(level, text): return ("h", level, text)
def P(text): return ("p", text)
def L(items): return ("ul", items)
def T(headers, rows): return ("table", headers, rows)
def NOTE(text): return ("note", text)


def _fmt(v):
    if isinstance(v, float):
        return f"{v:.4g}"
    if isinstance(v, (list, tuple)):
        return ", ".join(_fmt(x) for x in v)
    if isinstance(v, dict):
        return "; ".join(f"{k}={_fmt(x)}" for k, x in v.items())
    return "" if v is None else str(v)


def _aut(a: dict) -> str:
    if not a or not a.get("linked"):
        return (a or {}).get("note", "—")
    item = a["item"]
    pos = (f"{a['catalog_id']} termasuk daftar (strength {item['security_strength_bits']}, status {item['status_nist']})"
           if a["in_list"] else f"{a['catalog_id']} TIDAK ada dalam daftar")
    return (f"{a['total_combinations']} kombinasi ({_fmt(a['by_primitive'])}); {a['warnings_high']} peringatan tinggi; {pos} — "
            f"`{Path(a['file']).name}`")


def build_document(r: dict) -> list:
    pf, alg = r["profile"], r["profile"]["algorithm"]
    sel = r["selection"]
    D = [H(1, f"Dokumen Penetapan Metode & Parameter Pengujian — {alg['id']}"),
         P(f"**{r['app']['name']}** v{r['app']['version']} · Modul **{r['unit']['module']}** · "
           f"Unit {r['unit']['code']} — *{r['unit']['title']}* · {r['unit']['skkni']}"),
         P(f"Dibangkitkan: {r['generated_at']} · mode: {r['mode']} · objek: **{pf['product']['name']}** "
           f"({pf['product']['version']})"),
         NOTE("Seluruh data produk bersifat ILUSTRATIF. Nilai bertanda HASIL UJI LANGSUNG dihitung oleh kode; "
              "HASIL LITERATUR dikutip dari rujukan; PERLU_VERIFIKASI wajib dicek ke sumber primer."),
         H(2, "Ringkasan"),
         L([f"Kelas primitif: **{alg['primitive_label']}**, struktur {alg['structure']}",
            f"Tingkat akses penguji: **{pf['tester']['access']}**",
            f"KAT implementasi referensi: **{r['kat']['passed']}/{r['kat']['total']} {r['kat']['status']}**",
            f"Komponen berpotensi lemah: **{', '.join(r['component_summary']['berpotensi_lemah']) or '—'}**; "
            f"perlu perhatian: {', '.join(r['component_summary']['perhatian']) or '—'}",
            f"Metode terpilih: **{len(sel['selected'])}** · ditolak: {len(sel['rejected'])} · "
            f"upaya {sel['effort_hours_total']} / {sel['person_hours_available']} jam-orang ({sel['effort_fit']})"])]

    # Bab 1
    par = alg["parameters"]
    D += [H(2, "1. Profil Produk & Ruang Lingkup (KUK 1.1)"),
          T(["Atribut", "Nilai"], [
              ["Produk", f"{pf['product']['name']} {pf['product']['version']}"],
              ["Deskripsi", pf["product"]["description"]],
              ["Algoritma", alg["name"]], ["Kelas primitif", alg["primitive_label"]],
              ["Struktur", alg["structure"]], ["Mode", alg.get("mode") or "—"],
              ["Parameter", _fmt(par)],
              ["Platform", pf["implementation"].get("platform_desc", pf["implementation"]["platform"])],
              ["Bahasa / binding", f"{pf['implementation']['language']} / {_fmt(pf['implementation'].get('bindings', []))}"],
              ["Antarmuka", _fmt(pf["implementation"]["interfaces"])],
              ["RBG", pf["implementation"]["rng"]],
              ["Catatan implementasi", _fmt(alg.get("implementation_notes", {})) or "—"],
              ["Tingkat akses", f"{pf['tester']['access']} — {pf['tester']['access_desc']}"],
              ["Kelengkapan profil", pf["completeness"]["status"]],
              ["Daftar algoritma diuji (algo_catalog)", _aut(pf.get("algorithms_under_test", {}))]]),
          P("**Ruang lingkup:**"), L(pf["tester"]["scope"]),
          P("**Di luar ruang lingkup:**"), L(pf["tester"]["out_of_scope"] or ["—"])]

    # Bab 2
    D += [H(2, "2. Daftar Referensi Standar & Tren Serangan"),
          H(3, "2.1 Daftar Referensi Standar (KUK 1.3)"),
          T(["ID", "Organisasi", "Judul", "Tahun", "Peran"],
            [[x["id"], x["org"], x["title"] + (f" *({x['note']})*" if x["note"] else "") + (" `PERLU_VERIFIKASI`" if x["verify"] else ""),
              x["year"], x["role_label"]] for x in r["references"]]),
          H(3, "2.2 Tren Serangan Relevan (KUK 1.2) — HASIL LITERATUR"),
          P(f"{r['attacks']['summary']['total']} serangan relevan setelah filter profil; praktis: "
            f"{len(r['attacks']['summary']['practical'])}; perlu verifikasi: {len(r['attacks']['summary']['needs_verification'])}."),
          T(["ID", "Kategori", "Serangan", "Model", "Kompleksitas (log2 T/D/M)", "Ronde", "Status", "Rujukan"],
            [[a["id"], a["category"], a["name"], a["model"],
              "/".join("—" if a["complexity"].get(k) is None else f"{a['complexity'][k]:g}" for k in ("time_log2", "data_log2", "memory_log2")),
              a.get("rounds", ""), a["status"] + (" `PERLU_VERIFIKASI`" if a["verify"] == "PERLU_VERIFIKASI" else ""),
              "; ".join(a["refs"])] for a in r["attacks"]["items"]])]

    # Bab 3
    D += [H(2, "3. Hasil Telaah Komponen (KUK 2.1)"),
          T(["ID", "Komponen", "Lapis", "Status", "Tag kelemahan"],
            [[c["id"], c["name"], c["layer"], c["status"], ", ".join(c["tags"])] for c in r["components"]])]
    for c in r["components"]:
        D += [H(3, f"{c['id']} — {c['name']} [{c['status']}]"), P(c["role"]),
              T(["Pemeriksaan", "Nilai", "Acuan", "Status", "Sumber", "Catatan"],
                [[f["check"], _fmt(f["value"]), f["reference"], f["status"], f["source"], f["note"]] for f in c["findings"]])]

    # Bab 4
    D += [H(2, "4. Matriks Komponen–Kelemahan–Metode (KUK 2.2)"),
          P("Tiga lapis pengujian × level uji (metode yang dipetakan):"),
          T(["Lapis \\ Level", "Unit", "Integrasi", "Sistem"],
            [[l, *[", ".join(v[lv]) or "—" for lv in ("Unit", "Integrasi", "Sistem")]] for l, v in r["layer_summary"].items()]),
          T(["Komponen", "Status", "Kelemahan", "Serangan", "Metode", "Lapis", "Level", "Standar", "Akses"],
            [[f"{m['component_id']} {m['component']}", m["component_status"], m["weakness"], ", ".join(m["attacks"]) or "—",
              f"{m['method_id']} {m['method']}", m["layer"], m["level"], ", ".join(m["standards"]),
              ("✔ " if m["access_ok"] else "✘ ") + m["access_required"]] for m in r["matrix"]])]

    # Bab 5
    b = r["resources"]["benchmark"]
    cap = r["resources"]["capacity"]
    D += [H(2, "5. Estimasi Sumber Daya & Kelayakan (KUK 2.3)"),
          H(3, "5.1 Benchmark mesin lab — HASIL UJI LANGSUNG"),
          T(["Operasi", "ops/detik"], [[k, f"{v:,.1f}"] for k, v in b["ops_per_sec"].items()]),
          P(f"Mesin: {b['machine']['system']}, {b['machine']['processor']}, {b['machine']['logical_cpus']} CPU logis, "
            f"Python {b['machine']['python']}. Implementasi: {b['implementation']}."),
          H(3, "5.2 Sumber daya internal"),
          T(["Sumber daya", "Nilai"], [
              ["SDM", "; ".join(f"{x['count']}× {x['role']} ({', '.join(x['skills'])})" for x in cap["personnel"])],
              ["Jadwal", f"{cap['schedule_days']} hari × {cap['hours_per_day']} jam = {cap['person_hours']} jam-orang"],
              ["Komputasi", f"{cap['compute']['machines']} mesin × {cap['compute']['cores']} core, {cap['compute']['ram_gb']} GB RAM"],
              ["Anggaran komputasi per metode", f"{cap['method_budget_hours']} jam"],
              ["Faktor speedup", cap["speedup_factor"]],
              ["Tools tersedia", ", ".join(cap["tools_available"])]]),
          H(3, "5.3 Estimasi waktu per metode"),
          P("Waktu = 2^k ÷ (ops/detik × core × speedup)."),
          T(["Metode", "Operasi", "k (log2)", "Estimasi penuh", "Data", "Memori", "Varian tereduksi", "Status", "Alasan"],
            [[e["method_id"], e["op"], e["ops_log2"], e["time_human"], e["data"], e["memory"],
              (f"{e['reduced']['desc']} — {e['reduced']['time_human']}" if e.get("reduced") else "—"), e["status"], e["reason"]]
             for e in r["resources"]["estimates"]])]

    # Bab 6
    D += [H(2, "6. Metode Terpilih & Parameter Pengujian (KUK 3.1, 3.2)"),
          H(3, "6.1 Penetapan metode"), P(f"Rumus: {sel['formula']}; ambang ≥ {sel['threshold']}."),
          T(["Metode", "Lapis", "Level", "Varian", "Skor", "Target", "Alasan"],
            [[f"{s['method_id']} {s['name']}", s["layer"], s["level"], s["variant"], s["score"],
              ", ".join(s["weak_targets"] or s["targets"]), s["reason"]] for s in sel["selected"]]),
          P("**Metode yang ditolak:**"),
          T(["Metode", "Kelayakan", "Skor", "Alasan"],
            [[f"{s['method_id']} {s['name']}", s["feasibility"], s["score"], s["reason"]] for s in sel["rejected"]]),
          H(3, "6.2 Parameter pengujian")]
    for mid, prm in r["parameters"].items():
        rows = []
        for k, v in prm.items():
            if k == "tests" and isinstance(v, dict):
                v = "; ".join(f"{t} {_fmt(pp)}".strip() for t, pp in v.items())
            rows.append([k, _fmt(v)])
        D += [P(f"**{mid}**"), T(["Parameter", "Nilai"], rows)]

    # Bab 7
    D += [H(2, "7. Matriks Keterlacakan"),
          T(["ID", "Persyaratan", "Objek", "Metode uji", "Rujukan", "Kriteria lulus"],
            [[t["id"], t["requirement"], t["object"], f"{t['method_id']} ({t['variant']})", t["reference"], t["pass_criteria"]]
             for t in r["traceability"]])]

    # Bab 8
    D += [H(2, "8. Lampiran: Peta KUK → Bagian Dokumen → File Kode"),
          T(["KUK", "Deskripsi", "Bagian", "File kode", "Unit test"],
            [[k["kuk"], k["desc"], k["section"], ", ".join(k["files"]), ", ".join(k["tests"])] for k in r["kuk_map"]]),
          P(f"KAT rinci ({r['kat']['algorithm']}):"),
          T(["Vektor", "Sumber", "Hasil"], [[c["id"], c["source"], "LULUS" if c["pass"] else "GAGAL"] for c in r["kat"]["cases"]])]
    return D


# ---------------------------------------------------------------- renderer
def _md_cell(x):
    return _fmt(x).replace("|", "\\|").replace("\n", " ")


def to_markdown(doc: list) -> str:
    out = []
    for b in doc:
        if b[0] == "h":
            out.append("#" * b[1] + " " + b[2])
        elif b[0] == "p":
            out.append(b[1])
        elif b[0] == "note":
            out.append("> " + b[1])
        elif b[0] == "ul":
            out.append("\n".join(f"- {x}" for x in b[1]))
        elif b[0] == "table":
            out.append("| " + " | ".join(b[1]) + " |")
            out[-1] += "\n|" + "---|" * len(b[1])
            out[-1] += "".join("\n| " + " | ".join(_md_cell(c) for c in row) + " |" for row in b[2])
        out.append("")
    return "\n".join(out)


def to_docx(doc: list, path: Path) -> bool:
    try:
        import docx
        from docx.shared import Pt
    except ImportError:
        return False
    import re
    d = docx.Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(10)

    def runs(par, text):
        for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
            if part.startswith("**"):
                par.add_run(part[2:-2]).bold = True
            elif part.startswith("`"):
                par.add_run(part[1:-1]).italic = True
            else:
                par.add_run(part.replace("*", ""))
    for b in doc:
        if b[0] == "h":
            d.add_heading(b[2], level=min(b[1], 4) - (1 if b[1] > 1 else 0))
        elif b[0] in ("p", "note"):
            runs(d.add_paragraph(style="Intense Quote" if b[0] == "note" else None), b[1])
        elif b[0] == "ul":
            for x in b[1]:
                runs(d.add_paragraph(style="List Bullet"), x)
        elif b[0] == "table":
            t = d.add_table(rows=1, cols=len(b[1]))
            t.style = "Light Grid Accent 1"
            for i, h in enumerate(b[1]):
                t.rows[0].cells[i].text = h
            for row in b[2]:
                cells = t.add_row().cells
                for i, c in enumerate(row):
                    cells[i].text = _fmt(c).replace("**", "").replace("`", "")
            d.add_paragraph()
    d.save(path)
    return True


def write_all(result: dict, out_dir: Path, docx: bool = True) -> dict:
    out_dir.mkdir(parents=True, exist_ok=True)
    files = {}
    p = out_dir / JSON_NAME
    p.write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8")
    files["json"] = p
    doc = build_document(result)
    p = out_dir / MD_NAME
    p.write_text(to_markdown(doc), encoding="utf-8")
    files["md"] = p
    if docx and to_docx(doc, out_dir / DOCX_NAME):
        files["docx"] = out_dir / DOCX_NAME
    with open(out_dir / CSV_MAP, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["component_id", "component", "component_status", "weakness_tag", "weakness", "attacks",
                    "method_id", "method", "layer", "level", "standards", "access_required", "access_ok"])
        for m in result["matrix"]:
            w.writerow([m["component_id"], m["component"], m["component_status"], m["weakness_tag"], m["weakness"],
                        ";".join(m["attacks"]), m["method_id"], m["method"], m["layer"], m["level"],
                        ";".join(m["standards"]), m["access_required"], m["access_ok"]])
    files["csv_map"] = out_dir / CSV_MAP
    with open(out_dir / CSV_TRACE, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["id", "type", "requirement", "object", "method_id", "method", "variant", "reference", "pass_criteria"])
        for t in result["traceability"]:
            w.writerow([t[k] for k in ("id", "type", "requirement", "object", "method_id", "method", "variant", "reference", "pass_criteria")])
    files["csv_trace"] = out_dir / CSV_TRACE
    return files
